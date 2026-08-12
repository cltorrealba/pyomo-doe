from __future__ import annotations

import csv
import math
import tempfile
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import matplotlib.pyplot as plt
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


FERMENTATION_MODEL_DIR = Path(__file__).resolve().parents[1]
LOT_DIR = FERMENTATION_MODEL_DIR / "data" / "Laboratorio 2026" / "MBDoE_2026" / "DOE_Lote_2"
Y15_DIR = LOT_DIR / "Datos Y15 lote 2"
OUTPUT_XLSX = LOT_DIR / "DOE_Lote_2_ethanol_manual_entry.xlsx"
PREVIEW_DIR = Path(tempfile.gettempdir()) / "pyomo_doe_lot2_ethanol_preview"

ETHANOL_DENSITY_G_ML = 0.789
PERCENT_VV_TO_G_L = ETHANOL_DENSITY_G_ML * 10.0

PROCESS_META = {
    "F1": {
        "lab": "004",
        "design_name": "synthetic_glucose_rich_fructose_pulse",
    },
    "F2": {
        "lab": "005",
        "design_name": "synthetic_lit_SM410_24C_highN_strip",
    },
    "F3": {
        "lab": "006",
        "design_name": "synthetic_viable_biomass_step",
    },
}

ANALYTE_COLUMNS = {
    "GLUCOSE-320": "glucose_g_l",
    "GLU-FRU-320": "sugar_total_g_l",
    "AMMONIA": "ammonia_mg_l",
    "PAN": "pan_mg_l",
    "GLYCEROL": "glycerol_g_l",
    "ACETIC ACID": "acetic_acid_g_l",
    "PYRUVIC ACID": "pyruvic_acid_mg_l",
    "ACETALDEHID-1000": "acetaldehyde_mg_l",
}

HEADERS = [
    "lot_id",
    "process",
    "design_name",
    "sample_id",
    "sample_sequence",
    "sample_collection_datetime",
    "t_h",
    "time_source",
    "y15_data_available",
    "y15_analysis_datetime_start",
    "y15_analysis_datetime_end",
    "y15_source_files",
    "glucose_g_l",
    "fructose_g_l",
    "sugar_total_g_l",
    "ammonia_mg_l",
    "pan_mg_l",
    "yan_calc_mg_l_raw",
    "glycerol_g_l",
    "acetic_acid_g_l",
    "pyruvic_acid_mg_l",
    "acetaldehyde_mg_l",
    "ethanol_percent_vv_rep1",
    "ethanol_percent_vv_rep2",
    "ethanol_percent_vv_final",
    "ethanol_g_l_manual",
    "ethanol_g_l_for_model",
    "ethanol_measurement_datetime",
    "ethanol_method",
    "operator",
    "qc_flag",
    "notes",
]

INPUT_COLUMNS = {
    "sample_collection_datetime",
    "t_h",
    "time_source",
    "ethanol_percent_vv_rep1",
    "ethanol_percent_vv_rep2",
    "ethanol_g_l_manual",
    "ethanol_measurement_datetime",
    "ethanol_method",
    "operator",
    "qc_flag",
    "notes",
}

FORMULA_COLUMNS = {
    "fructose_g_l",
    "yan_calc_mg_l_raw",
    "ethanol_percent_vv_final",
    "ethanol_g_l_for_model",
}


def parse_decimal(value: str) -> float | None:
    text = value.strip().replace(" ", "").replace(",", ".")
    if not text:
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    return number if math.isfinite(number) else None


def read_y15() -> dict[str, dict[str, object]]:
    if not Y15_DIR.exists():
        raise FileNotFoundError(f"Missing Y15 directory: {Y15_DIR}")

    exact_records: set[tuple[str, str, str, str, datetime, str]] = set()
    by_key: dict[tuple[str, str], list[tuple[float | None, str, datetime, str]]] = defaultdict(list)

    for path in sorted(Y15_DIR.glob("*.txt")):
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            for fields in csv.reader(handle, delimiter="\t"):
                if len(fields) < 6:
                    continue
                sample_id, analyte, _method, value_text, unit, datetime_text = fields[:6]
                sample_id = sample_id.strip()
                analyte = analyte.strip()
                timestamp = datetime.strptime(datetime_text.strip(), "%d-%m-%Y %H:%M:%S")
                exact_key = (
                    sample_id,
                    analyte,
                    value_text.strip(),
                    unit.strip(),
                    timestamp,
                    path.name,
                )
                if exact_key in exact_records:
                    continue
                exact_records.add(exact_key)
                by_key[(sample_id, analyte)].append(
                    (parse_decimal(value_text), unit.strip(), timestamp, path.name)
                )

    conflicts = {key: values for key, values in by_key.items() if len(values) > 1}
    if conflicts:
        conflict_names = ", ".join(f"{sample}/{analyte}" for sample, analyte in conflicts)
        raise ValueError(f"Conflicting Y15 records require review: {conflict_names}")

    samples: dict[str, dict[str, object]] = defaultdict(
        lambda: {"timestamps": [], "source_files": set(), "analytes": {}}
    )
    for (sample_id, analyte), values in by_key.items():
        value, unit, timestamp, source_file = values[0]
        samples[sample_id]["timestamps"].append(timestamp)
        samples[sample_id]["source_files"].add(source_file)
        samples[sample_id]["analytes"][analyte] = {"value": value, "unit": unit}
    return dict(samples)


def build_rows(y15: dict[str, dict[str, object]]) -> dict[str, list[dict[str, object]]]:
    rows_by_process: dict[str, list[dict[str, object]]] = {}
    for process, meta in PROCESS_META.items():
        process_rows = []
        for sequence in range(1, 19):
            sample_id = f"DOE-LAB{meta['lab']}-{sequence}"
            sample = y15.get(sample_id)
            analytes = sample["analytes"] if sample else {}
            timestamps = sorted(sample["timestamps"]) if sample else []
            row: dict[str, object] = {
                "lot_id": "DOE_Lote_2",
                "process": process,
                "design_name": meta["design_name"],
                "sample_id": sample_id,
                "sample_sequence": sequence,
                "sample_collection_datetime": None,
                "t_h": None,
                "time_source": "pending_manual_entry",
                "y15_data_available": bool(sample),
                "y15_analysis_datetime_start": timestamps[0] if timestamps else None,
                "y15_analysis_datetime_end": timestamps[-1] if timestamps else None,
                "y15_source_files": "; ".join(sorted(sample["source_files"])) if sample else "",
            }
            for analyte, column in ANALYTE_COLUMNS.items():
                row[column] = analytes.get(analyte, {}).get("value")
            process_rows.append(row)
        rows_by_process[process] = process_rows
    return rows_by_process


def style_sheet(ws, *, editable: bool) -> None:
    ws.showGridLines = False
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.zoomScale = 80

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True, size=10)
    input_fill = PatternFill("solid", fgColor="FFF2CC")
    formula_fill = PatternFill("solid", fgColor="E2F0D9")
    source_fill = PatternFill("solid", fgColor="DDEBF7")
    readonly_fill = PatternFill("solid", fgColor="E7E6E6")
    missing_fill = PatternFill("solid", fgColor="FCE4D6")
    thin_gray = Side(style="thin", color="D9E2F3")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color="17365D"))
    ws.row_dimensions[1].height = 42

    header_to_col = {cell.value: cell.column for cell in ws[1]}
    for name, col_idx in header_to_col.items():
        fill = source_fill
        if editable and name in INPUT_COLUMNS:
            fill = input_fill
        elif name in FORMULA_COLUMNS:
            fill = formula_fill
        elif not editable:
            fill = readonly_fill
        for row_idx in range(2, ws.max_row + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.fill = fill
            cell.border = Border(bottom=thin_gray)
            cell.alignment = Alignment(vertical="center", wrap_text=name in {"design_name", "notes"})

    y15_col = header_to_col["y15_data_available"]
    y15_letter = ws.cell(row=1, column=y15_col).column_letter
    ws.conditional_formatting.add(
        f"A2:AF{ws.max_row}",
        FormulaRule(formula=[f"${y15_letter}2=FALSE"], fill=missing_fill),
    )

    widths = {
        "A": 15, "B": 9, "C": 42, "D": 18, "E": 12, "F": 22, "G": 10, "H": 23,
        "I": 16, "J": 22, "K": 22, "L": 34, "M": 13, "N": 13, "O": 15, "P": 14,
        "Q": 12, "R": 16, "S": 13, "T": 14, "U": 17, "V": 18, "W": 18, "X": 18,
        "Y": 18, "Z": 16, "AA": 18, "AB": 24, "AC": 18, "AD": 16, "AE": 16, "AF": 34,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 24
    for column in ("F", "J", "K", "AB"):
        for row_idx in range(2, ws.max_row + 1):
            ws[f"{column}{row_idx}"].number_format = "yyyy-mm-dd hh:mm"
    for column in ("G", "M", "N", "O", "P", "Q", "R", "S", "T", "U", "V", "W", "X", "Y", "Z", "AA"):
        for row_idx in range(2, ws.max_row + 1):
            ws[f"{column}{row_idx}"].number_format = "0.00"

    if editable:
        qc_validation = DataValidation(
            type="list",
            formula1='"OK,not_measured,below_LOQ,above_range,suspect,rerun,discard"',
            allow_blank=True,
        )
        ws.add_data_validation(qc_validation)
        qc_validation.add(f"AE2:AE{ws.max_row}")


def write_process_sheet(ws, rows: list[dict[str, object]], table_name: str) -> None:
    ws.append(HEADERS)
    for row_data in rows:
        ws.append([row_data.get(name) for name in HEADERS])

    for row_idx in range(2, ws.max_row + 1):
        ws[f"N{row_idx}"] = f'=IF(OR(M{row_idx}="",O{row_idx}=""),"",O{row_idx}-M{row_idx})'
        ws[f"R{row_idx}"] = f'=IF(COUNTA(P{row_idx}:Q{row_idx})=0,"",SUM(P{row_idx}:Q{row_idx}))'
        ws[f"Y{row_idx}"] = (
            f'=IF(COUNTA(W{row_idx}:X{row_idx})=0,"",AVERAGE(W{row_idx}:X{row_idx}))'
        )
        ws[f"AA{row_idx}"] = (
            f'=IF(Z{row_idx}<>"",Z{row_idx},IF(Y{row_idx}<>"",Y{row_idx}*{PERCENT_VV_TO_G_L:.2f},""))'
        )

    table = Table(displayName=table_name, ref=f"A1:AF{ws.max_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=False,
        showColumnStripes=False,
    )
    ws.add_table(table)
    style_sheet(ws, editable=True)


def write_all_samples_sheet(ws) -> None:
    ws.append(HEADERS)
    target_row = 2
    for process in PROCESS_META:
        for source_row in range(2, 20):
            for column_idx in range(1, len(HEADERS) + 1):
                letter = ws.cell(row=1, column=column_idx).column_letter
                ws.cell(target_row, column_idx, f"='{process}'!{letter}{source_row}")
            target_row += 1
    table = Table(displayName="Lot2AllSamples", ref=f"A1:AF{ws.max_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=False,
        showColumnStripes=False,
    )
    ws.add_table(table)
    style_sheet(ws, editable=False)


def build_workbook(rows_by_process: dict[str, list[dict[str, object]]]) -> Workbook:
    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    readme.showGridLines = False
    readme.sheet_view.zoomScale = 95

    readme_rows = [
        ["DOE Lote 2 - manual ethanol entry"],
        [],
        ["Purpose", "Enter ethanol measurements for DOE-LAB004, DOE-LAB005 and DOE-LAB006."],
        ["Where to enter", "Use only sheets F1, F2 and F3. All_samples is a formula-linked read-only view."],
        ["Yellow cells", "Manual inputs. Enter % v/v replicates or a direct g/L override."],
        ["Green cells", "Formula-derived values. Do not overwrite."],
        ["Blue cells", "Context parsed from the Y15 session files."],
        ["Orange rows", "Expected sample ID with no corresponding Y15 record; alcohol data may still be entered."],
        ["% v/v conversion", f"ethanol_g_l_for_model = ethanol_percent_vv_final x {PERCENT_VV_TO_G_L:.2f}, unless ethanol_g_l_manual is supplied."],
        ["Sample time", "Y15 timestamps are analysis times, not sample collection times. Fill sample_collection_datetime and/or t_h when known."],
        ["Missing expected IDs", "F1: 7, 14, 17; F2: 7, 14, 17; F3: 7, 14, 17, 18."],
        ["QC flags", "OK, not_measured, below_LOQ, above_range, suspect, rerun, discard."],
        ["Source", str(Y15_DIR)],
    ]
    for row in readme_rows:
        readme.append(row)
    readme["A1"].font = Font(bold=True, color="FFFFFF", size=16)
    readme["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    readme["A1"].alignment = Alignment(vertical="center")
    readme.merge_cells("A1:B1")
    readme.row_dimensions[1].height = 30
    readme.column_dimensions["A"].width = 24
    readme.column_dimensions["B"].width = 115
    for row_idx in range(3, readme.max_row + 1):
        readme[f"A{row_idx}"].font = Font(bold=True, color="1F4E78")
        readme[f"A{row_idx}"].alignment = Alignment(vertical="top")
        readme[f"B{row_idx}"].alignment = Alignment(wrap_text=True, vertical="top")
        readme.row_dimensions[row_idx].height = 30

    for process, rows in rows_by_process.items():
        ws = wb.create_sheet(process)
        write_process_sheet(ws, rows, f"Lot2{process}Ethanol")

    all_samples = wb.create_sheet("All_samples")
    write_all_samples_sheet(all_samples)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    return wb


def verify_workbook(path: Path) -> dict[str, object]:
    wb = load_workbook(path, data_only=False)
    expected_sheets = ["README", "F1", "F2", "F3", "All_samples"]
    if wb.sheetnames != expected_sheets:
        raise AssertionError(f"Unexpected sheets: {wb.sheetnames}")
    checks: dict[str, object] = {"sheets": wb.sheetnames}
    for process in PROCESS_META:
        ws = wb[process]
        if ws.max_row != 19 or ws.max_column != len(HEADERS):
            raise AssertionError(f"Unexpected shape for {process}: {ws.max_row}x{ws.max_column}")
        for row_idx in range(2, 20):
            for column in ("N", "R", "Y", "AA"):
                if not str(ws[f"{column}{row_idx}"].value).startswith("="):
                    raise AssertionError(f"Missing formula in {process}!{column}{row_idx}")
        if ws["W2"].fill.fgColor.rgb[-6:] != "FFF2CC":
            raise AssertionError(f"Manual-input styling missing in {process}!W2")
        if ws["Y2"].fill.fgColor.rgb[-6:] != "E2F0D9":
            raise AssertionError(f"Formula styling missing in {process}!Y2")
        if len(ws.data_validations.dataValidation) != 1:
            raise AssertionError(f"QC validation missing in {process}")
        checks[process] = {
            "rows": ws.max_row - 1,
            "y15_available": sum(bool(ws[f"I{row_idx}"].value) for row_idx in range(2, 20)),
            "formulas": sum(
                str(ws[f"{column}{row_idx}"].value).startswith("=")
                for row_idx in range(2, 20)
                for column in ("N", "R", "Y", "AA")
            ),
        }
    all_samples = wb["All_samples"]
    if all_samples.max_row != 55 or all_samples.max_column != len(HEADERS):
        raise AssertionError(
            f"Unexpected All_samples shape: {all_samples.max_row}x{all_samples.max_column}"
        )
    formula_errors = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and any(
                    marker in value for marker in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")
                ):
                    formula_errors.append(f"{ws.title}!{cell.coordinate}: {value}")
    if formula_errors:
        raise AssertionError(f"Formula errors found: {formula_errors[:5]}")
    checks["All_samples"] = {"rows": all_samples.max_row - 1, "linked_columns": len(HEADERS)}
    checks["formula_errors"] = 0
    return checks


def render_previews(path: Path) -> list[Path]:
    wb = load_workbook(path, data_only=False)
    PREVIEW_DIR.mkdir(parents=True, exist_ok=True)
    previews = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        max_rows = min(ws.max_row, 20)
        max_cols = min(ws.max_column, 12 if sheet_name != "README" else 2)
        data = [
            [ws.cell(row=row, column=column).value for column in range(1, max_cols + 1)]
            for row in range(1, max_rows + 1)
        ]
        text = [["" if value is None else str(value)[:55] for value in row] for row in data]
        fig_width = 16 if sheet_name != "README" else 14
        fig_height = max(3.5, 0.42 * max_rows)
        fig, ax = plt.subplots(figsize=(fig_width, fig_height))
        ax.axis("off")
        table = ax.table(cellText=text, cellLoc="left", loc="upper left", bbox=[0, 0, 1, 1])
        table.auto_set_font_size(False)
        table.set_fontsize(6.5 if sheet_name != "README" else 8)
        for (row, _column), cell in table.get_celld().items():
            cell.set_edgecolor("#D9E2F3")
            if row == 0:
                cell.set_facecolor("#1F4E78")
                cell.get_text().set_color("white")
                cell.get_text().set_weight("bold")
        ax.set_title(f"{path.name} - {sheet_name}", fontsize=12, pad=8)
        preview = PREVIEW_DIR / f"{sheet_name}.png"
        fig.savefig(preview, dpi=160, bbox_inches="tight")
        plt.close(fig)
        previews.append(preview)

        if sheet_name in PROCESS_META:
            first_column = 22  # V: acetaldehyde plus all ethanol-entry columns through AF.
            last_column = 32
            entry_data = [
                [ws.cell(row=row, column=column).value for column in range(first_column, last_column + 1)]
                for row in range(1, max_rows + 1)
            ]
            entry_text = [
                ["" if value is None else str(value)[:55] for value in row]
                for row in entry_data
            ]
            fig, ax = plt.subplots(figsize=(18, fig_height))
            ax.axis("off")
            table = ax.table(cellText=entry_text, cellLoc="left", loc="upper left", bbox=[0, 0, 1, 1])
            table.auto_set_font_size(False)
            table.set_fontsize(7)
            for (row, column), cell in table.get_celld().items():
                cell.set_edgecolor("#D9E2F3")
                if row == 0:
                    cell.set_facecolor("#1F4E78")
                    cell.get_text().set_color("white")
                    cell.get_text().set_weight("bold")
                else:
                    absolute_column = first_column + column
                    header = HEADERS[absolute_column - 1]
                    if header in INPUT_COLUMNS:
                        cell.set_facecolor("#FFF2CC")
                    elif header in FORMULA_COLUMNS:
                        cell.set_facecolor("#E2F0D9")
                    else:
                        cell.set_facecolor("#DDEBF7")
            ax.set_title(f"{path.name} - {sheet_name} alcohol entry", fontsize=12, pad=8)
            entry_preview = PREVIEW_DIR / f"{sheet_name}_alcohol_entry.png"
            fig.savefig(entry_preview, dpi=160, bbox_inches="tight")
            plt.close(fig)
            previews.append(entry_preview)
    return previews


def main() -> None:
    y15 = read_y15()
    rows_by_process = build_rows(y15)
    workbook = build_workbook(rows_by_process)
    workbook.save(OUTPUT_XLSX)
    checks = verify_workbook(OUTPUT_XLSX)
    previews = render_previews(OUTPUT_XLSX)
    print(f"Wrote: {OUTPUT_XLSX}")
    print(f"Checks: {checks}")
    print("Previews:")
    for preview in previews:
        print(preview)


if __name__ == "__main__":
    main()
