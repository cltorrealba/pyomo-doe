from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


LABORATORY_DIR = Path(__file__).resolve().parent
RESULTS_DIR = LABORATORY_DIR / "results" / "lot1_data_preview"
PROCESSED_DIR = RESULTS_DIR / "processed"
OUT_DIR = RESULTS_DIR / "manual_entry_templates"
OUT_DIR.mkdir(parents=True, exist_ok=True)

OUT_XLSX = OUT_DIR / "DOE_Lote_1_ethanol_manual_entry.xlsx"
OUT_CSV = OUT_DIR / "DOE_Lote_1_ethanol_manual_entry.csv"

ETHANOL_DENSITY_G_ML = 0.789
PERCENT_VV_TO_G_L = ETHANOL_DENSITY_G_ML * 10.0

PROCESS_META = {
    "F1": "synthetic_lit_SM410_18C_highN_ester",
    "F2": "synthetic_high_biomass_low_N_maintenance",
    "F3": "synthetic_fructose_rich_glucose",
}


def load_schedule() -> pd.DataFrame:
    y15 = pd.read_csv(PROCESSED_DIR / "lot1_y15_wide_processed.csv")
    oculyze = pd.read_csv(PROCESSED_DIR / "lot1_oculyze_processed.csv")

    y15["sample_datetime"] = pd.to_datetime(y15["sample_datetime"], errors="coerce")
    oculyze["sample_datetime"] = pd.to_datetime(oculyze["sample_datetime"], errors="coerce")

    density_cols = [
        "sample_id_clean",
        "Density",
        "x_total_g_l",
        "x_viable_g_l",
        "x_dead_g_l",
        "Viability",
        "Budding Index",
    ]
    density = oculyze[density_cols].rename(columns={"sample_id_clean": "sample_id"})

    df = y15.merge(density, on="sample_id", how="left")
    df["lot_id"] = "DOE_Lote_1"
    df["design_name"] = df["process"].map(PROCESS_META)

    cols = [
        "lot_id",
        "process",
        "design_name",
        "sample_id",
        "sample_datetime",
        "t_h",
        "time_source",
        "Density",
        "sugar_total_g_l",
        "glucose_g_l",
        "fructose_g_l",
        "yan_calc_mg_l_censored",
        "x_viable_g_l",
        "ethanol_percent_vv_rep1",
        "ethanol_percent_vv_rep2",
        "ethanol_percent_vv_final",
        "ethanol_g_l_manual",
        "ethanol_g_l_for_model",
        "ethanol_measurement_datetime",
        "ethanol_method",
        "operator",
        "qc_flag",
        "notes",
    ]
    for col in cols:
        if col not in df.columns:
            df[col] = ""

    df = df[cols].sort_values(["process", "sample_datetime", "sample_id"]).reset_index(drop=True)
    df["sample_datetime"] = df["sample_datetime"].dt.strftime("%Y-%m-%d %H:%M")
    return df


def write_dataframe(ws, df: pd.DataFrame):
    headers = list(df.columns)
    ws.append(headers)
    for row in df.itertuples(index=False):
        ws.append(list(row))

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    widths = {
        "A": 14,
        "B": 10,
        "C": 42,
        "D": 18,
        "E": 18,
        "F": 10,
        "G": 26,
        "H": 12,
        "I": 14,
        "J": 12,
        "K": 12,
        "L": 14,
        "M": 12,
        "N": 18,
        "O": 18,
        "P": 18,
        "Q": 16,
        "R": 18,
        "S": 24,
        "T": 18,
        "U": 16,
        "V": 14,
        "W": 30,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    header_to_col = {cell.value: cell.column for cell in ws[1]}
    rep1_col = get_column_letter(header_to_col["ethanol_percent_vv_rep1"])
    rep2_col = get_column_letter(header_to_col["ethanol_percent_vv_rep2"])
    final_col = get_column_letter(header_to_col["ethanol_percent_vv_final"])
    manual_col = get_column_letter(header_to_col["ethanol_g_l_manual"])
    model_col = get_column_letter(header_to_col["ethanol_g_l_for_model"])
    qc_col = get_column_letter(header_to_col["qc_flag"])

    input_fill = PatternFill("solid", fgColor="FFF2CC")
    formula_fill = PatternFill("solid", fgColor="E2F0D9")
    input_headers = [
        "ethanol_percent_vv_rep1",
        "ethanol_percent_vv_rep2",
        "ethanol_g_l_manual",
        "ethanol_measurement_datetime",
        "ethanol_method",
        "operator",
        "qc_flag",
        "notes",
    ]
    formula_headers = ["ethanol_percent_vv_final", "ethanol_g_l_for_model"]
    for name in input_headers:
        col = get_column_letter(header_to_col[name])
        for row in range(1, ws.max_row + 1):
            ws[f"{col}{row}"].fill = input_fill
    for name in formula_headers:
        col = get_column_letter(header_to_col[name])
        for row in range(1, ws.max_row + 1):
            ws[f"{col}{row}"].fill = formula_fill

    for row in range(2, ws.max_row + 1):
        ws[f"{final_col}{row}"] = (
            f'=IF(COUNTA({rep1_col}{row}:{rep2_col}{row})=0,"",'
            f"AVERAGE({rep1_col}{row}:{rep2_col}{row}))"
        )
        ws[f"{model_col}{row}"] = (
            f'=IF({manual_col}{row}<>"",{manual_col}{row},'
            f'IF({final_col}{row}<>"",{final_col}{row}*{PERCENT_VV_TO_G_L:.2f},""))'
        )

    qc_validation = DataValidation(
        type="list",
        formula1='"OK,not_measured,below_LOQ,above_range,suspect,rerun,discard"',
        allow_blank=True,
    )
    ws.add_data_validation(qc_validation)
    qc_validation.add(f"{qc_col}2:{qc_col}{ws.max_row}")

    numeric_cols = [
        "t_h",
        "Density",
        "sugar_total_g_l",
        "glucose_g_l",
        "fructose_g_l",
        "yan_calc_mg_l_censored",
        "x_viable_g_l",
        "ethanol_percent_vv_rep1",
        "ethanol_percent_vv_rep2",
        "ethanol_percent_vv_final",
        "ethanol_g_l_manual",
        "ethanol_g_l_for_model",
    ]
    for name in numeric_cols:
        if name in header_to_col:
            col = get_column_letter(header_to_col[name])
            for row in range(2, ws.max_row + 1):
                ws[f"{col}{row}"].number_format = "0.00"


def build_workbook(df: pd.DataFrame):
    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    readme.append(["DOE Lote 1 ethanol manual entry"])
    readme.append([])
    readme.append(["Fill ethanol_percent_vv_rep1 and optionally ethanol_percent_vv_rep2 if the instrument reports % v/v."])
    readme.append(["ethanol_percent_vv_final is calculated as the average of available % v/v replicates."])
    readme.append([f"ethanol_g_l_for_model uses ethanol_g_l_manual if entered; otherwise it uses % v/v * {PERCENT_VV_TO_G_L:.2f}."])
    readme.append(["Do not overwrite formula columns unless you intentionally want a manual override."])
    readme.append(["Recommended qc_flag values: OK, not_measured, below_LOQ, above_range, suspect, rerun, discard."])
    readme.append(["Time zero is the first Oculyze process timestamp for each fermentation."])
    readme["A1"].font = Font(bold=True, size=14)
    readme.column_dimensions["A"].width = 120

    ws_all = wb.create_sheet("All_samples")
    write_dataframe(ws_all, df)

    for process in ["F1", "F2", "F3"]:
        ws = wb.create_sheet(process)
        write_dataframe(ws, df[df["process"].eq(process)].copy())

    return wb


def main():
    df = load_schedule()
    df.to_csv(OUT_CSV, index=False)
    wb = build_workbook(df)
    wb.save(OUT_XLSX)
    print(f"Wrote {OUT_XLSX}")
    print(f"Wrote {OUT_CSV}")


if __name__ == "__main__":
    main()
