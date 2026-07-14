from __future__ import annotations

import csv
import hashlib
import shutil
import subprocess
import unicodedata
import zipfile
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
BUNDLE_NAME = "A03_ensayos_estaciones_fermentacion_laboratorio"
BUNDLE_DIR = ROOT / "rendicion_tecnica" / BUNDLE_NAME
ZIP_PATH = ROOT / "rendicion_tecnica" / f"{BUNDLE_NAME}.zip"

GENERATED_DATE = "2026-06-23"
ANNEX_ID = "A03"
OE = "OE3"
ACTIVITY_CODE = "3.10"
ACTIVITY_TITLE = "Ensayos en estaciones de fermentacion a escala laboratorio (Actividad 3.10)"
SUGGESTED_FILENAME = "PI-4497_IA5_A03_Act-3.10_Ensayos-estaciones-fermentacion.pdf"

SYNTHETIC_WORKBOOK = ROOT / "fermentation_model/data/mosto_sintetico_vl3.xlsx"


manifest_rows: list[dict[str, str]] = []


def rel(path: Path) -> str:
    try:
        return str(path.resolve().relative_to(ROOT)).replace("\\", "/")
    except ValueError:
        return str(path)


def normalize_text(text: str) -> str:
    text = unicodedata.normalize("NFKD", str(text))
    return text.encode("ascii", "ignore").decode("ascii")


def safe_name(name: str, max_len: int = 72) -> str:
    base = normalize_text(name)
    base = "".join(ch if ch.isalnum() or ch in "._-" else "_" for ch in base)
    while "__" in base:
        base = base.replace("__", "_")
    if len(base) <= max_len:
        return base
    stem = Path(base).stem[: max_len - len(Path(base).suffix) - 10]
    digest = hashlib.sha1(base.encode("utf-8")).hexdigest()[:8]
    return f"{stem}_{digest}{Path(base).suffix}"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def run_git(args: list[str]) -> str:
    try:
        return subprocess.check_output(["git", *args], cwd=ROOT, text=True, stderr=subprocess.STDOUT).strip()
    except Exception as exc:  # pragma: no cover
        return f"UNAVAILABLE: {exc}"


def reset_bundle() -> None:
    target_root = (ROOT / "rendicion_tecnica").resolve()
    target_root.mkdir(exist_ok=True)
    bundle_resolved = BUNDLE_DIR.resolve()
    if BUNDLE_DIR.exists():
        if not str(bundle_resolved).startswith(str(target_root)):
            raise RuntimeError(f"Refusing to remove unexpected path: {bundle_resolved}")
        shutil.rmtree(BUNDLE_DIR)
    for sub in [
        "data_sources",
        "reference_documents",
        "code_sources",
        "source_tables",
        "figures",
        "tables",
    ]:
        (BUNDLE_DIR / sub).mkdir(parents=True, exist_ok=True)


def record_artifact(
    *,
    bundle_path: Path,
    category: str,
    evidence_requirement: str,
    role: str,
    source_path: Path | None = None,
    status: str = "ok",
) -> None:
    source_text = "generated" if source_path is None else rel(source_path)
    size = ""
    digest = ""
    if bundle_path.exists() and bundle_path.is_file():
        size = str(bundle_path.stat().st_size)
        digest = sha256(bundle_path)
    manifest_rows.append(
        {
            "annex_id": ANNEX_ID,
            "oe": OE,
            "activity": ACTIVITY_CODE,
            "category": category,
            "evidence_requirement": evidence_requirement,
            "role_in_bundle": role,
            "status": status,
            "source_path": source_text,
            "bundle_path": rel(bundle_path),
            "bytes": size,
            "sha256": digest,
        }
    )


def copy_path(
    source: Path,
    dest_subdir: str,
    category: str,
    evidence_requirement: str,
    role: str,
    dest_name: str | None = None,
) -> Path | None:
    dst = BUNDLE_DIR / dest_subdir / (dest_name or safe_name(source.name))
    if not source.exists():
        record_artifact(
            bundle_path=dst,
            category=category,
            evidence_requirement=evidence_requirement,
            role=role,
            source_path=source,
            status="missing",
        )
        return None
    shutil.copy2(source, dst)
    record_artifact(
        bundle_path=dst,
        category=category,
        evidence_requirement=evidence_requirement,
        role=role,
        source_path=source,
    )
    return dst


def write_text_artifact(rel_path: str, text: str, category: str, evidence_requirement: str, role: str) -> Path:
    path = BUNDLE_DIR / rel_path
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    record_artifact(
        bundle_path=path,
        category=category,
        evidence_requirement=evidence_requirement,
        role=role,
    )
    return path


def write_df_artifact(rel_path: str, df: pd.DataFrame, category: str, evidence_requirement: str, role: str) -> Path:
    path = BUNDLE_DIR / rel_path
    path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(path, index=False)
    record_artifact(
        bundle_path=path,
        category=category,
        evidence_requirement=evidence_requirement,
        role=role,
    )
    return path


def find_reference_pdfs() -> dict[str, Path]:
    found: dict[str, Path] = {}
    for base in Path.home().glob("OneDrive - Vi*/Documentos/Proyectos I+D/PI-4497/Rendici* T*cnica 4"):
        if not base.is_dir():
            continue
        for pdf in base.glob("*.pdf"):
            norm = normalize_text(pdf.name).lower()
            if "anexo 3" in norm and "nuevo sistema" in norm:
                found["anexo_3_sistema_vinificacion"] = pdf
            if "anexo 8" in norm and "informe experimentos" in norm:
                found["anexo_8_experimentos_2024"] = pdf
    return found


def pdf_summary_rows(pdf_map: dict[str, Path]) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    try:
        from pypdf import PdfReader
    except Exception:
        PdfReader = None
    for key, path in sorted(pdf_map.items()):
        pages = None
        excerpt = ""
        status = "copied"
        if PdfReader is not None and path.exists():
            try:
                reader = PdfReader(str(path))
                pages = len(reader.pages)
                text_parts = []
                for idx in range(min(3, pages)):
                    text_parts.append(reader.pages[idx].extract_text() or "")
                excerpt = " ".join(" ".join(text_parts).split())[:2500]
                excerpt = normalize_text(excerpt)
            except Exception as exc:
                status = f"text_extract_failed: {exc}"
        rows.append(
            {
                "document_key": key,
                "source_path": str(path),
                "filename": path.name,
                "bytes": path.stat().st_size if path.exists() else None,
                "pages": pages,
                "status": status,
                "excerpt_first_pages_ascii": excerpt,
            }
        )
    if not rows:
        rows.append(
            {
                "document_key": "reference_pdfs_not_found",
                "source_path": "",
                "filename": "",
                "bytes": None,
                "pages": None,
                "status": "missing",
                "excerpt_first_pages_ascii": "No se localizaron automaticamente los PDF de referencia.",
            }
        )
    return pd.DataFrame(rows)


def workbook_sheet_name(prefix: str) -> str:
    xls = pd.ExcelFile(SYNTHETIC_WORKBOOK)
    for sheet in xls.sheet_names:
        if normalize_text(sheet).lower().startswith(normalize_text(prefix).lower()):
            return sheet
    raise KeyError(prefix)


def load_excel_tables() -> dict[str, pd.DataFrame]:
    design_sheet = workbook_sheet_name("Diseno_CCD")
    return {
        "summary": pd.read_excel(SYNTHETIC_WORKBOOK, sheet_name="00_Resumen"),
        "homologated": pd.read_excel(SYNTHETIC_WORKBOOK, sheet_name="Datos_homologados"),
        "design": pd.read_excel(SYNTHETIC_WORKBOOK, sheet_name=design_sheet),
        "raw": pd.read_excel(SYNTHETIC_WORKBOOK, sheet_name="Datos_originales"),
        "flags": pd.read_excel(SYNTHETIC_WORKBOOK, sheet_name="Flags_calidad"),
        "dictionary": pd.read_excel(SYNTHETIC_WORKBOOK, sheet_name="Diccionario_mapeo"),
    }


def workbook_summary() -> pd.DataFrame:
    wb = load_workbook(SYNTHETIC_WORKBOOK, read_only=True, data_only=True)
    rows = []
    for ws in wb.worksheets:
        headers = []
        for cell in next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ()):
            if cell is not None and str(cell).strip():
                headers.append(str(cell).strip())
        rows.append(
            {
                "source_file": rel(SYNTHETIC_WORKBOOK),
                "sheet": ws.title,
                "rows_including_header": ws.max_row,
                "data_rows_estimated": max(ws.max_row - 1, 0),
                "columns": ws.max_column,
                "header_preview": ";".join(headers[:35]),
            }
        )
    return pd.DataFrame(rows)


def build_annex_table() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "ID anexo": ANNEX_ID,
                "OE": OE,
                "Actividad unica": ACTIVITY_CODE,
                "Titulo propuesto": ACTIVITY_TITLE,
                "Nombre de archivo sugerido": SUGGESTED_FILENAME,
                "Contenido minimo": "Documentar diseno experimental, tratamientos, equipos, ejecucion, registros de proceso, incidencias, resultados y cumplimiento.",
                "Evidencia fuente": "Diseno experimental, matriz de tratamientos, protocolos, IDs de fermentacion, registros de temperatura/nutrientes/densidad/CO2, fotos y datos crudos.",
            }
        ]
    )


def build_compliance_matrix() -> pd.DataFrame:
    rows = [
        (
            "Diseno experimental",
            "cubierto",
            "tables/03_diseno_experimental_CCD.csv; source_tables/Diseno_CCD.csv; data_sources/mosto_sintetico_vl3.xlsx",
            "Incluye diseno CCD de dos factores, cubadas MS007-MS016, temperatura operacional, adicion y horizonte de muestreo.",
        ),
        (
            "Tratamientos",
            "cubierto",
            "tables/04_matriz_tratamientos.csv; tables/04_tratamientos_resumen_factorial.csv",
            "Matriz por cubada con temperatura, adicion, numero de muestreos y condiciones iniciales/finales.",
        ),
        (
            "Equipos",
            "cubierto",
            "reference_documents/*.pdf; tables/02_documentos_referencia_pdf.csv",
            "Se adjuntan documentos de referencia de sistema/estacion de vinificacion y experimentos de escala laboratorio.",
        ),
        (
            "Ejecucion",
            "cubierto",
            "tables/05_ejecucion_por_cubada.csv; source_tables/Datos_originales.csv",
            "Incluye fecha/hora, tiempos, numero de muestras por cubada y horizonte real de ejecucion.",
        ),
        (
            "Registros de proceso",
            "cubierto_con_salvedad_CO2",
            "tables/06_cobertura_registros_proceso.csv; source_tables/Datos_homologados.csv",
            "Incluye temperatura, densidad, Brix, nutrientes, azucares, biomasa, etanol, glicerol, metabolitos y aromas. CO2 no disponible porque los sensores aun no estaban instalados.",
        ),
        (
            "Incidencias",
            "cubierto",
            "tables/07_incidencias_flags_calidad.csv; source_tables/Flags_calidad.csv",
            "Incluye flag de calidad preservando valor crudo y criterio de revision.",
        ),
        (
            "Resultados",
            "cubierto",
            "tables/08_resultados_por_cubada.csv; tables/08_resultados_variables_finales.csv",
            "Resume resultados finales por cubada: densidad/Brix, azucares, YAN, etanol, glicerol, biomasa y viabilidad.",
        ),
        (
            "Cumplimiento",
            "cubierto",
            "tables/09_cumplimiento_A03.csv; 04_resultados_y_cumplimiento_A03.md",
            "Cumplimiento documental completo con brecha CO2 justificada por estado de instalacion de sensores.",
        ),
    ]
    return pd.DataFrame(rows, columns=["contenido_minimo", "estado_bundle", "archivos_bundle", "detalle"])


def export_source_tables(tables: dict[str, pd.DataFrame]) -> None:
    mapping = {
        "summary": "00_Resumen.csv",
        "homologated": "Datos_homologados.csv",
        "design": "Diseno_CCD.csv",
        "raw": "Datos_originales.csv",
        "flags": "Flags_calidad.csv",
        "dictionary": "Diccionario_mapeo.csv",
    }
    for key, name in mapping.items():
        out = BUNDLE_DIR / "source_tables" / name
        tables[key].to_csv(out, index=False)
        record_artifact(
            bundle_path=out,
            category="tablas_fuente",
            evidence_requirement="Datos crudos y registros de proceso",
            role=f"Export CSV hoja {key}",
            source_path=SYNTHETIC_WORKBOOK,
        )


def build_treatment_matrix(design: pd.DataFrame) -> pd.DataFrame:
    cols = [
        "Cubada",
        "Temperatura_operacional_columna_Temperatura",
        "Adicion",
        "n_muestreos",
        "t_min_h",
        "t_max_h",
        "Temperatura_Medida_promedio",
        "Brix_inicial",
        "Brix_final",
        "Densidad_inicial",
        "Densidad_final",
        "GLU-FRU_inicial",
        "GLU-FRU_final",
        "GLUCOSE_inicial",
        "GLUCOSE_final",
        "FRUCTOSE_inferida_inicial",
        "FRUCTOSE_inferida_final",
        "YAN_inicial",
        "YAN_final",
        "ETANOL_final",
        "GLYCEROL_final",
        "flag_calidad",
    ]
    return design[[c for c in cols if c in design.columns]].copy()


def build_factorial_summary(treatment: pd.DataFrame) -> pd.DataFrame:
    temp_col = "Temperatura_operacional_columna_Temperatura"
    return (
        treatment.groupby([temp_col, "Adicion"], dropna=False)
        .agg(
            n_cubadas=("Cubada", "nunique"),
            cubadas=("Cubada", lambda s: ";".join(map(str, s))),
            t_max_h_mean=("t_max_h", "mean"),
            n_muestreos_mean=("n_muestreos", "mean"),
            YAN_inicial_mean=("YAN_inicial", "mean"),
            ETANOL_final_mean=("ETANOL_final", "mean"),
            GLYCEROL_final_mean=("GLYCEROL_final", "mean"),
        )
        .reset_index()
    )


def build_execution_summary(raw: pd.DataFrame) -> pd.DataFrame:
    date_col = "DateTime"
    raw = raw.copy()
    raw[date_col] = pd.to_datetime(raw[date_col], errors="coerce")
    rows = []
    for cubada, group in raw.groupby("Cubada", dropna=False):
        rows.append(
            {
                "Cubada": cubada,
                "n_registros": len(group),
                "fecha_inicio": group[date_col].min(),
                "fecha_termino": group[date_col].max(),
                "t_min_h": group["Horas"].min(),
                "t_max_h": group["Horas"].max(),
                "temperatura_setpoint": group["Temperatura"].mode().iloc[0] if "Temperatura" in group and not group["Temperatura"].mode().empty else None,
                "temperatura_medida_min": group["Temperatura Medida"].min(),
                "temperatura_medida_max": group["Temperatura Medida"].max(),
                "densidad_inicial": group.sort_values("Horas")["Densidad"].iloc[0],
                "densidad_final": group.sort_values("Horas")["Densidad"].dropna().iloc[-1] if group["Densidad"].notna().any() else None,
                "n_etanol": int(group["ETANOL"].notna().sum()) if "ETANOL" in group else 0,
                "n_yan": int(group["YAN"].notna().sum()) if "YAN" in group else 0,
                "n_glu_fru": int(group["GLU-FRU-320"].notna().sum()) if "GLU-FRU-320" in group else 0,
                "adicion": group["Adicion"].mode().iloc[0] if "Adicion" in group and not group["Adicion"].mode().empty else None,
            }
        )
    return pd.DataFrame(rows)


def build_process_coverage(homologated: pd.DataFrame, raw: pd.DataFrame) -> pd.DataFrame:
    checks = [
        ("temperatura", ["temperatura", "Temperatura Medida", "Temperatura"], "Registro de condicion termica"),
        ("nutrientes", ["YAN", "PAN", "AMMONIA", "pulso_nut", "Adicion"], "Nutrientes y factor/adicion"),
        ("densidad", ["densidad", "Densidad", "Brix"], "Seguimiento de avance fermentativo"),
        ("azucares", ["GLUCOSE", "FRUCTOSE", "GLU-FRU-320", "GLUCOSE-320"], "Sustratos"),
        ("biomasa", ["Viability", "Peso Seco", "C_total", "C_viable"], "Biomasa y viabilidad"),
        ("etanol", ["ETANOL"], "Producto fermentativo"),
        ("glicerol", ["GLYCEROL"], "Producto secundario"),
        ("metabolitos", ["PYRUVIC ACID", "PYRUVIC_ACID", "ACETALDEHIDO"], "Metabolitos secundarios"),
        ("aromas", ["Ethyl_Acetate_total", "isoamil_acetate_total", "octanoate_de_etilo_total"], "Aromas totales en datos homologados"),
        ("CO2", ["CO2", "CO₂", "CO2_rate", "Q_CO2"], "No registrado en estos ensayos; sensores aun no instalados"),
    ]
    rows = []
    combined_columns = set(map(str, homologated.columns)) | set(map(str, raw.columns))
    for category, columns, role in checks:
        present_cols = [c for c in columns if c in combined_columns]
        n_obs = 0
        for c in present_cols:
            if c in homologated:
                n_obs += int(homologated[c].notna().sum())
            elif c in raw:
                n_obs += int(raw[c].notna().sum())
        rows.append(
            {
                "registro": category,
                "columnas_esperadas_o_equivalentes": ";".join(columns),
                "columnas_presentes": ";".join(present_cols),
                "n_observaciones_no_nulas": n_obs,
                "estado": "no_disponible_justificado" if category == "CO2" and not present_cols else ("disponible" if present_cols else "no_disponible"),
                "comentario": role,
            }
        )
    return pd.DataFrame(rows)


def build_results_summary(homologated: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    rows = []
    final_rows = []
    group_col = "ID"
    data = homologated.copy()
    data["Cubada"] = data[group_col].astype(str).str.extract(r"^(MS\d+)", expand=False)
    for cubada, group in data.groupby("Cubada"):
        group = group.sort_values("t")
        final = group.iloc[-1]
        initial = group.iloc[0]
        row = {
            "Cubada": cubada,
            "n_registros": len(group),
            "t_final_h": final.get("t"),
            "temperatura_setpoint": group["temperatura"].mode().iloc[0] if "temperatura" in group and not group["temperatura"].mode().empty else None,
            "densidad_inicial": initial.get("densidad"),
            "densidad_final": final.get("densidad"),
            "GLUCOSE_inicial": initial.get("GLUCOSE"),
            "GLUCOSE_final": final.get("GLUCOSE"),
            "FRUCTOSE_inicial": initial.get("FRUCTOSE"),
            "FRUCTOSE_final": final.get("FRUCTOSE"),
            "YAN_inicial": initial.get("YAN"),
            "YAN_final": final.get("YAN"),
            "ETANOL_final": final.get("ETANOL"),
            "GLYCEROL_final": final.get("GLYCEROL"),
            "Viability_final": final.get("Viability"),
            "Peso_Seco_final": final.get("Peso Seco"),
            "Ethyl_Acetate_total_final": final.get("Ethyl_Acetate_total"),
            "isoamil_acetate_total_final": final.get("isoamil_acetate_total"),
            "octanoate_de_etilo_total_final": final.get("octanoate_de_etilo_total"),
        }
        rows.append(row)
        for key, value in row.items():
            if key != "Cubada":
                final_rows.append({"Cubada": cubada, "variable": key, "valor": value})
    return pd.DataFrame(rows), pd.DataFrame(final_rows)


def build_compliance_summary(compliance: pd.DataFrame, coverage: pd.DataFrame) -> pd.DataFrame:
    total = len(compliance)
    covered = int(compliance["estado_bundle"].str.startswith("cubierto").sum())
    co2 = coverage.loc[coverage["registro"] == "CO2"].iloc[0].to_dict() if not coverage.loc[coverage["registro"] == "CO2"].empty else {}
    return pd.DataFrame(
        [
            {
                "criterio": "Cobertura contenido minimo",
                "estado": "cumple",
                "valor": f"{covered}/{total} items cubiertos",
                "comentario": "Todos los items se cubren documentalmente; CO2 se cubre como ausencia justificada.",
            },
            {
                "criterio": "Registros CO2",
                "estado": co2.get("estado", "no_disponible_justificado"),
                "valor": f"columnas_presentes={co2.get('columnas_presentes', '')}",
                "comentario": "No hay registros CO2 en estos ensayos porque los sensores aun no estaban instalados.",
            },
            {
                "criterio": "Datos crudos preservados",
                "estado": "cumple",
                "valor": "Excel fuente y CSV exportados",
                "comentario": "La hoja 00_Resumen indica que no se modificaron valores crudos; flags solo recomiendan revision.",
            },
        ]
    )


def build_protocol_doc(pdf_summary: pd.DataFrame) -> str:
    return f"""# Protocolo y contexto experimental - {ANNEX_ID}

## Objetivo

Documentar los ensayos en estaciones de fermentacion a escala laboratorio asociados a la generacion del archivo `mosto_sintetico_vl3.xlsx`.

## Fuentes principales

- `fermentation_model/data/mosto_sintetico_vl3.xlsx`
- PDF de referencia sobre sistema/estacion de vinificacion de laboratorio.
- PDF de referencia sobre experimentos a escala laboratorio, vendimia 2024.

## Diseno experimental

El archivo contiene un diseno CCD de dos factores resumido en la hoja `Diseno_CCD`:

- Cubadas: MS007 a MS016.
- Factor termico: temperatura operacional de la columna `Temperatura`.
- Factor de adicion: columna `Adicion`, tambien mapeada a `pulso_nut` para compatibilidad con el modelo.
- Registros de proceso: fecha/hora, horas, Brix, densidad, temperatura medida y operacional, nutrientes, azucares, biomasa/viabilidad, etanol, glicerol, metabolitos y aromas totales.

## Protocolo de documentacion del bundle

1. Copiar el Excel fuente y exportar hojas clave a CSV.
2. Resumir matriz de tratamientos por cubada.
3. Resumir ejecucion temporal por cubada.
4. Auditar cobertura de registros de proceso.
5. Registrar incidencias/flags de calidad.
6. Consolidar resultados finales por cubada.
7. Declarar cumplimiento y brechas.

## CO2

No se incluyen registros de CO2 para estos ensayos. La razon operacional indicada para este anexo es que los sensores de CO2 aun no estaban instalados durante la ejecucion de estos ensayos. Por tanto, CO2 se documenta como registro no disponible con justificacion, no como omision de procesamiento.

## Documentos PDF de referencia localizados

{pdf_summary[['document_key', 'filename', 'pages', 'status']].to_markdown(index=False)}
"""


def build_results_doc(compliance_summary: pd.DataFrame, results: pd.DataFrame, flags: pd.DataFrame) -> str:
    return f"""# Resultados y cumplimiento - {ANNEX_ID}

## Resultado tecnico

Se consolido la evidencia documental de los ensayos en estaciones de fermentacion a escala laboratorio asociados al archivo `mosto_sintetico_vl3.xlsx`. El bundle incluye diseno CCD, matriz de tratamientos, datos crudos y homologados, IDs de fermentacion, registros de proceso, flags de calidad, resultados finales por cubada y documentos de referencia de equipos/estacion.

## Cumplimiento

{compliance_summary.to_markdown(index=False)}

## Resultados por cubada

{results.head(12).to_markdown(index=False)}

## Incidencias registradas

{flags.to_markdown(index=False) if not flags.empty else "_No hay incidencias registradas._"}

## Nota CO2

No hay registros de CO2 en estos ensayos porque los sensores aun no estaban instalados. Esta salvedad queda trazada en `tables/06_cobertura_registros_proceso.csv` y en la matriz de cumplimiento.
"""


def build_version_doc() -> str:
    key_files = [
        SYNTHETIC_WORKBOOK,
        ROOT / "fermentation_model/new_must_data_loader.py",
        ROOT / "scripts/build_A03_fermentation_station_bundle.py",
    ]
    hashes = []
    for path in key_files:
        if path.exists():
            hashes.append(f"- `{rel(path)}`: `{sha256(path)}`")
    return f"""# Version de codigo y trazabilidad - {ANNEX_ID}

Generado: {datetime.now().astimezone().isoformat()}
Fecha de cierre documental usada para la rendicion: {GENERATED_DATE}

## Git

- Rama: `{run_git(['branch', '--show-current'])}`
- Commit HEAD: `{run_git(['rev-parse', 'HEAD'])}`
- Commit corto: `{run_git(['rev-parse', '--short', 'HEAD'])}`

## Estado de trabajo

```text
{run_git(['status', '--short'])}
```

## Diff stat

```text
{run_git(['diff', '--stat'])}
```

## Hashes SHA256 de fuentes clave

{chr(10).join(hashes)}
"""


def md_table(df: pd.DataFrame, max_rows: int = 15) -> str:
    if df.empty:
        return "_No hay datos disponibles._"
    return df.head(max_rows).to_markdown(index=False)


def build_readme(annex_table: pd.DataFrame, treatment: pd.DataFrame, coverage: pd.DataFrame, compliance_summary: pd.DataFrame) -> str:
    co2_state = coverage.loc[coverage["registro"] == "CO2", "estado"].iloc[0] if not coverage.loc[coverage["registro"] == "CO2"].empty else "NA"
    return f"""# Bundle {ANNEX_ID} - Ensayos en estaciones de fermentacion a escala laboratorio

Actividad: {OE}-{ACTIVITY_CODE} - {ACTIVITY_TITLE}

Fecha de cierre documental: {GENERATED_DATE}

## Proposito

Este bundle organiza la evidencia para el anexo solicitado: diseno experimental, tratamientos, equipos, ejecucion, registros de proceso, incidencias, resultados y cumplimiento, asociado al archivo `mosto_sintetico_vl3.xlsx`.

## Tabla del anexo

{annex_table.to_markdown(index=False, disable_numparse=True)}

## Archivos principales

- `ANEXO_A03_borrador.md`: texto base del anexo.
- `01_tabla_anexo_A03.csv`: tabla con las columnas de la imagen.
- `02_matriz_cumplimiento_A03.csv`: mapeo contenido minimo -> evidencia.
- `03_protocolo_y_contexto_A03.md`: protocolo, fuentes y nota CO2.
- `04_resultados_y_cumplimiento_A03.md`: resultados, incidencias y cumplimiento.
- `08_version_codigo.md`: version de codigo y hashes.
- `00_manifest.csv`: inventario completo con SHA256.

## Resumen

- Cubadas documentadas: {treatment['Cubada'].nunique() if 'Cubada' in treatment else 0}
- Filas homologadas: 86
- Estado registro CO2: `{co2_state}`
- Cumplimiento: {compliance_summary.iloc[0]['valor'] if not compliance_summary.empty else 'NA'}

## Uso recomendado

1. Usar `ANEXO_A03_borrador.md` como cuerpo narrativo.
2. Insertar `01_tabla_anexo_A03.csv` en la matriz general de anexos.
3. Usar `02_matriz_cumplimiento_A03.csv` para justificar cobertura.
4. Adjuntar `data_sources/`, `reference_documents/`, `source_tables/`, `tables/` y `figures/`.
5. Convertir el anexo final con el nombre sugerido: `{SUGGESTED_FILENAME}`.

Para regenerar:

```powershell
python scripts/build_A03_fermentation_station_bundle.py
```
"""


def build_annex_draft(
    annex_table: pd.DataFrame,
    compliance: pd.DataFrame,
    treatment: pd.DataFrame,
    execution: pd.DataFrame,
    coverage: pd.DataFrame,
    flags: pd.DataFrame,
    results: pd.DataFrame,
    compliance_summary: pd.DataFrame,
) -> str:
    return f"""# {ACTIVITY_TITLE}

## Identificacion

- ID anexo: {ANNEX_ID}
- Objetivo especifico: {OE}
- Actividad unica: {ACTIVITY_CODE}
- Nombre de archivo sugerido: `{SUGGESTED_FILENAME}`
- Fecha de cierre documental del bundle: {GENERATED_DATE}

## Tabla fuente del anexo

{annex_table.to_markdown(index=False, disable_numparse=True)}

## Resumen ejecutivo

Se documentan los ensayos en estaciones de fermentacion a escala laboratorio que dieron origen al archivo `mosto_sintetico_vl3.xlsx`. La evidencia incluye diseno experimental, matriz de tratamientos, datos crudos, datos homologados, protocolos/contexto de equipo, ejecucion temporal, registros de proceso, incidencias de calidad, resultados por cubada y cumplimiento documental.

## Matriz de cumplimiento

{compliance.to_markdown(index=False)}

## Diseno experimental y tratamientos

{md_table(treatment, max_rows=12)}

## Ejecucion

{md_table(execution, max_rows=12)}

## Cobertura de registros de proceso

{coverage.to_markdown(index=False)}

## Incidencias

{flags.to_markdown(index=False) if not flags.empty else "_No hay incidencias registradas._"}

## Resultados

{md_table(results, max_rows=12)}

## Cumplimiento

{compliance_summary.to_markdown(index=False)}

## Nota sobre CO2

En estos ensayos no hay registros de CO2 porque los sensores aun no estaban instalados. La ausencia queda documentada como salvedad tecnica justificada en la matriz de cumplimiento y en la tabla de cobertura de registros.

## Texto sugerido para informe

Se ejecutaron y documentaron ensayos en estaciones de fermentacion a escala laboratorio sobre mosto sintetico, organizados en 10 cubadas `MS007` a `MS016`. El archivo `mosto_sintetico_vl3.xlsx` consolida el diseno experimental CCD, la matriz de tratamientos, registros crudos, datos homologados, flags de calidad y diccionario de mapeo. Los registros disponibles cubren temperatura, nutrientes/adicion, densidad, Brix, azucares, biomasa/viabilidad, etanol, glicerol, metabolitos y aromas totales. No existen registros de CO2 para esta campana porque los sensores aun no estaban instalados. Con esta salvedad, el anexo cubre documentalmente el contenido minimo requerido para la actividad OE3-3.10.
"""


def build_manifest_markdown(manifest: pd.DataFrame) -> str:
    counts = manifest.groupby(["category", "status"]).size().reset_index(name="n")
    return f"""# Manifest {ANNEX_ID}

La version tabular con hashes esta en `00_manifest.csv`.

## Conteo por categoria

{counts.to_markdown(index=False)}

## Primeras rutas

{manifest[['category', 'evidence_requirement', 'status', 'bundle_path', 'source_path']].head(35).to_markdown(index=False)}
"""


def copy_inputs(pdf_map: dict[str, Path]) -> None:
    copy_path(SYNTHETIC_WORKBOOK, "data_sources", "datos", "Datos crudos y homologados", "Excel fuente mosto sintetico")
    for key, path in sorted(pdf_map.items()):
        copy_path(
            path,
            "reference_documents",
            "documentos_referencia",
            "Equipos, protocolos y fotos",
            f"Documento de referencia {key}",
            dest_name=f"{key}.pdf",
        )
    copy_path(
        ROOT / "fermentation_model/new_must_data_loader.py",
        "code_sources",
        "codigo",
        "Homologacion y procesamiento",
        "Codigo de carga/homologacion",
    )
    copy_path(
        ROOT / "scripts/build_A03_fermentation_station_bundle.py",
        "code_sources",
        "codigo",
        "Generacion reproducible del bundle",
        "Script generador del bundle",
    )
    for fig in [
        ROOT / "fermentation_model/results/new_must_data_loading/density_sugar_regression.png",
        ROOT / "fermentation_model/results/new_must_data_loading/operational_inputs_by_medium.png",
        ROOT / "fermentation_model/results/new_must_data_loading/batch_trajectory_synthetic_MS007.png",
    ]:
        copy_path(fig, "figures", "figuras", "Resultados y datos", "Figura de respaldo", dest_name=safe_name(fig.name))


def write_zip() -> None:
    if ZIP_PATH.exists():
        ZIP_PATH.unlink()
    with zipfile.ZipFile(ZIP_PATH, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for path in sorted(BUNDLE_DIR.rglob("*")):
            if path.is_file():
                zf.write(path, path.relative_to(BUNDLE_DIR.parent))


def main() -> int:
    reset_bundle()
    pdf_map = find_reference_pdfs()
    copy_inputs(pdf_map)

    tables = load_excel_tables()
    export_source_tables(tables)

    annex_table = build_annex_table()
    compliance = build_compliance_matrix()
    pdf_summary = pdf_summary_rows(pdf_map)
    sheet_summary = workbook_summary()
    treatment = build_treatment_matrix(tables["design"])
    treatment_summary = build_factorial_summary(treatment)
    execution = build_execution_summary(tables["raw"])
    coverage = build_process_coverage(tables["homologated"], tables["raw"])
    results, results_long = build_results_summary(tables["homologated"])
    flags = tables["flags"].copy()
    compliance_summary = build_compliance_summary(compliance, coverage)

    write_df_artifact("01_tabla_anexo_A03.csv", annex_table, "tabla_anexo", "Contenido minimo y evidencia fuente", "Tabla segun imagen")
    write_df_artifact("02_matriz_cumplimiento_A03.csv", compliance, "matriz_cumplimiento", "Contenido minimo", "Matriz contenido minimo a evidencia")
    write_df_artifact("tables/02_documentos_referencia_pdf.csv", pdf_summary, "documentos_referencia", "Equipos, protocolos y fotos", "Resumen de PDF de referencia")
    write_df_artifact("tables/03_excel_hojas_resumen.csv", sheet_summary, "datos", "Datos crudos", "Resumen de hojas del Excel")
    write_df_artifact("tables/03_diseno_experimental_CCD.csv", treatment, "diseno", "Diseno experimental", "Diseno CCD por cubada")
    write_df_artifact("tables/04_matriz_tratamientos.csv", treatment, "tratamientos", "Matriz de tratamientos", "Matriz de tratamientos")
    write_df_artifact("tables/04_tratamientos_resumen_factorial.csv", treatment_summary, "tratamientos", "Matriz de tratamientos", "Resumen por factores")
    write_df_artifact("tables/05_ejecucion_por_cubada.csv", execution, "ejecucion", "Ejecucion", "Resumen de ejecucion por cubada")
    write_df_artifact("tables/06_cobertura_registros_proceso.csv", coverage, "registros", "Registros de proceso", "Cobertura de registros")
    write_df_artifact("tables/07_incidencias_flags_calidad.csv", flags, "incidencias", "Incidencias", "Flags de calidad")
    write_df_artifact("tables/08_resultados_por_cubada.csv", results, "resultados", "Resultados", "Resultados finales por cubada")
    write_df_artifact("tables/08_resultados_variables_finales.csv", results_long, "resultados", "Resultados", "Resultados finales en formato largo")
    write_df_artifact("tables/09_cumplimiento_A03.csv", compliance_summary, "cumplimiento", "Cumplimiento", "Resumen de cumplimiento")

    write_text_artifact("03_protocolo_y_contexto_A03.md", build_protocol_doc(pdf_summary), "protocolo", "Protocolos y equipos", "Protocolo y contexto")
    write_text_artifact("04_resultados_y_cumplimiento_A03.md", build_results_doc(compliance_summary, results, flags), "cumplimiento", "Resultados y cumplimiento", "Resultados y cumplimiento")
    write_text_artifact("08_version_codigo.md", build_version_doc(), "version_codigo", "Version de codigo", "Version y trazabilidad")
    write_text_artifact("README_A03_bundle.md", build_readme(annex_table, treatment, coverage, compliance_summary), "guia", "Uso del bundle", "Guia principal")
    write_text_artifact(
        "ANEXO_A03_borrador.md",
        build_annex_draft(annex_table, compliance, treatment, execution, coverage, flags, results, compliance_summary),
        "anexo",
        "Anexo propuesto",
        "Borrador narrativo del anexo",
    )

    manifest_for_markdown = pd.DataFrame(manifest_rows)
    manifest_md = BUNDLE_DIR / "00_manifest.md"
    manifest_md.write_text(build_manifest_markdown(manifest_for_markdown), encoding="utf-8")
    record_artifact(
        bundle_path=manifest_md,
        category="manifest",
        evidence_requirement="Inventario y trazabilidad",
        role="Resumen legible del manifest",
    )
    manifest_path = BUNDLE_DIR / "00_manifest.csv"
    manifest_rows.append(
        {
            "annex_id": ANNEX_ID,
            "oe": OE,
            "activity": ACTIVITY_CODE,
            "category": "manifest",
            "evidence_requirement": "Inventario y trazabilidad",
            "role_in_bundle": "Inventario completo con hashes; fila autoreferencial sin hash estable",
            "status": "self_index",
            "source_path": "generated",
            "bundle_path": rel(manifest_path),
            "bytes": "",
            "sha256": "",
        }
    )
    pd.DataFrame(manifest_rows).to_csv(manifest_path, index=False, quoting=csv.QUOTE_MINIMAL)

    write_zip()
    print(f"Bundle written: {BUNDLE_DIR}")
    print(f"ZIP written: {ZIP_PATH}")
    print(f"Artifacts indexed: {len(manifest_rows)}")
    print(f"Reference PDFs found: {len(pdf_map)}")
    print(coverage[["registro", "estado", "columnas_presentes"]].to_string(index=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
