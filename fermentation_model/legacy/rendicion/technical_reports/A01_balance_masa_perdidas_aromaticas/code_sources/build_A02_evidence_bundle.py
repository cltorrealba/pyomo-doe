from __future__ import annotations

import csv
import hashlib
import json
import shutil
import subprocess
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Iterable

import numpy as np
import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
BUNDLE_NAME = "A02_balance_masa_perdidas_aromaticas"
BUNDLE_DIR = ROOT / "rendicion_tecnica" / BUNDLE_NAME
ZIP_PATH = ROOT / "rendicion_tecnica" / f"{BUNDLE_NAME}.zip"

GENERATED_DATE = "2026-06-22"
ACTIVITY_CODE = "OE2-2.13"
ACTIVITY_NAME = "Desarrollar y calibrar modelos de balance de masa para cuantificar perdidas aromaticas"

CONTEXT_ALIASES = {
    "new_must_glycerol_estimability_doe": "nmg_doe",
    "secondary_fit_capacity": "sec_fit",
    "secondary_joint_campaign_doe": "sec_joint",
    "aroma_campaign_doe": "aroma",
    "aroma_joint_campaign_doe": "aroma_joint",
}


def rel(path: Path) -> str:
    try:
        return str(path.resolve().relative_to(ROOT)).replace("\\", "/")
    except ValueError:
        return str(path)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def run_git(args: list[str]) -> str:
    try:
        out = subprocess.check_output(["git", *args], cwd=ROOT, text=True, stderr=subprocess.STDOUT)
        return out.strip()
    except Exception as exc:  # pragma: no cover - used for evidence capture only
        return f"UNAVAILABLE: {exc}"


def reset_bundle() -> None:
    target_root = (ROOT / "rendicion_tecnica").resolve()
    target_root.mkdir(exist_ok=True)
    bundle_resolved = BUNDLE_DIR.resolve()
    if BUNDLE_DIR.exists():
        if not str(bundle_resolved).startswith(str(target_root)):
            raise RuntimeError(f"Refusing to remove unexpected path: {bundle_resolved}")
        shutil.rmtree(BUNDLE_DIR)
    BUNDLE_DIR.mkdir(parents=True, exist_ok=True)
    for sub in [
        "data_sources",
        "code_sources",
        "notebooks",
        "source_reports",
        "source_tables",
        "figures",
        "tables",
    ]:
        (BUNDLE_DIR / sub).mkdir(parents=True, exist_ok=True)


manifest_rows: list[dict[str, str]] = []


def record_artifact(
    *,
    bundle_path: Path,
    category: str,
    evidence_requirement: str,
    role: str,
    source_path: Path | None = None,
    status: str = "ok",
) -> None:
    source_text = "generated" if source_path is None else rel(source_path)
    sha = ""
    size = ""
    if bundle_path.exists() and bundle_path.is_file():
        sha = sha256(bundle_path)
        size = str(bundle_path.stat().st_size)
    manifest_rows.append(
        {
            "activity": ACTIVITY_CODE,
            "category": category,
            "evidence_requirement": evidence_requirement,
            "role_in_bundle": role,
            "status": status,
            "source_path": source_text,
            "bundle_path": rel(bundle_path),
            "bytes": size,
            "sha256": sha,
        }
    )


def copy_file(
    source: str,
    dest_subdir: str,
    category: str,
    evidence_requirement: str,
    role: str,
    namespace: bool = False,
) -> Path | None:
    src = ROOT / source
    dst_name = contextual_name(src) if namespace else src.name
    dst = BUNDLE_DIR / dest_subdir / dst_name
    if not src.exists():
        record_artifact(
            bundle_path=dst,
            category=category,
            evidence_requirement=evidence_requirement,
            role=role,
            source_path=src,
            status="missing",
        )
        return None
    shutil.copy2(src, dst)
    record_artifact(
        bundle_path=dst,
        category=category,
        evidence_requirement=evidence_requirement,
        role=role,
        source_path=src,
    )
    return dst


def copy_glob(
    pattern: str,
    dest_subdir: str,
    category: str,
    evidence_requirement: str,
    role_prefix: str,
) -> None:
    for src in sorted(ROOT.glob(pattern)):
        if not src.is_file():
            continue
        dst = BUNDLE_DIR / dest_subdir / contextual_name(src)
        shutil.copy2(src, dst)
        record_artifact(
            bundle_path=dst,
            category=category,
            evidence_requirement=evidence_requirement,
            role=f"{role_prefix}: {src.name}",
            source_path=src,
        )


def contextual_name(src: Path) -> str:
    """Return a stable bundle filename that avoids collisions across result dirs."""
    if src.parent.name == "plots":
        context = src.parent.parent.name
        prefix = f"{CONTEXT_ALIASES.get(context, context)}__plots"
    else:
        context = src.parent.name
        prefix = CONTEXT_ALIASES.get(context, context)
    name = f"{prefix}__{src.name}"
    if len(name) <= 62:
        return name
    digest = hashlib.sha1(rel(src).encode("utf-8")).hexdigest()[:8]
    suffix = src.suffix
    available = max(12, 62 - len(prefix) - len(digest) - len(suffix) - 4)
    return f"{prefix}__{src.stem[:available]}__{digest}{suffix}"


def write_text_artifact(
    rel_path: str,
    text: str,
    category: str,
    evidence_requirement: str,
    role: str,
) -> Path:
    path = BUNDLE_DIR / rel_path
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    record_artifact(
        bundle_path=path,
        category=category,
        evidence_requirement=evidence_requirement,
        role=role,
    )
    return path


def write_df_artifact(
    rel_path: str,
    df: pd.DataFrame,
    category: str,
    evidence_requirement: str,
    role: str,
) -> Path:
    path = BUNDLE_DIR / rel_path
    path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(path, index=False)
    record_artifact(
        bundle_path=path,
        category=category,
        evidence_requirement=evidence_requirement,
        role=role,
    )
    return path


def find_line(path: str, needle: str) -> int | None:
    full = ROOT / path
    if not full.exists():
        return None
    for idx, line in enumerate(full.read_text(encoding="utf-8", errors="replace").splitlines(), start=1):
        if needle in line:
            return idx
    return None


def summarize_workbook(path: Path) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    if not path.exists():
        return rows
    wb = load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        header_values = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        headers = [str(c).strip() for c in header_values if c is not None and str(c).strip()]
        aroma_cols = [
            h
            for h in headers
            if "total" in h.lower()
            and any(key in h.lower() for key in ["acetate", "octano", "benz", "hexil", "isoamil", "phenyl", "aroma"])
        ]
        key_cols = [
            h
            for h in headers
            if h
            in {
                "ID",
                "fecha_hora",
                "t",
                "temperatura",
                "densidad",
                "Viability",
                "Peso Seco",
                "GLUCOSE",
                "FRUCTOSE",
                "PAN",
                "AMMONIA",
                "YAN",
                "GLYCEROL",
                "PYRUVIC ACID",
                "ACETALDEHIDO",
                "ETANOL",
                "pulso_nut",
            }
        ]
        rows.append(
            {
                "source_file": rel(path),
                "sheet": ws.title,
                "rows_including_header": ws.max_row,
                "data_rows_estimated": max(ws.max_row - 1, 0),
                "columns": ws.max_column,
                "key_process_columns": ";".join(key_cols),
                "aroma_total_columns": ";".join(aroma_cols),
                "header_preview": ";".join(headers[:30]),
            }
        )
    return rows


def build_data_summary() -> pd.DataFrame:
    data_files = [
        ROOT / "fermentation_model/data/Calibration_data_vl3.xlsx",
        ROOT / "fermentation_model/data/mosto_sintetico_vl3.xlsx",
        ROOT / "fermentation_model/data/mosto_natural_xthiol.xlsx",
    ]
    rows: list[dict[str, object]] = []
    for path in data_files:
        rows.extend(summarize_workbook(path))
    return pd.DataFrame(rows)


def build_variable_dictionary() -> pd.DataFrame:
    rows = [
        ("X", "Biomasa viable", "Peso Seco, Viability, C_viable", "Estado dinamico calibrado"),
        ("Xd", "Biomasa muerta", "Viability, C_total/C_viable cuando disponible", "Estado extendido para mortalidad"),
        ("N", "Nitrogeno asimilable", "YAN, PAN, AMMONIA", "Balance de consumo de nitrogeno"),
        ("G", "Glucosa", "GLUCOSE", "Balance de consumo de azucar"),
        ("F", "Fructosa", "FRUCTOSE o FRUCTOSE_inferida", "Balance de consumo de azucar"),
        ("E", "Etanol", "ETANOL, Cf Alcolyzer", "Producto fermentativo calibrado"),
        ("Gly", "Glicerol", "GLYCEROL", "Producto secundario calibrado"),
        ("Pyr", "Acido piruvico", "PYRUVIC ACID", "Estado secundario de diagnostico"),
        ("AcAld", "Acetaldehido", "ACETALDEHIDO", "Estado secundario de diagnostico"),
        ("A_liq", "Aroma en fase liquida", "aroma total en liquido", "Estado de aroma en el modelo"),
        ("A_loss", "Aroma perdido por gas", "prediccion de stripping", "Estado acumulado de perdida"),
        ("A_cond", "Aroma terminal en condensado", "prediccion por eficiencia de trampa", "Cierre de masa aromatico"),
        ("Q_CO2", "Flujo volumetrico de CO2", "CO2_rate modelado", "Forzante de stripping"),
        ("K_lg", "Coeficiente gas/liquido", "UNIFAC/surrogado", "Particion para perdida aromatica"),
    ]
    return pd.DataFrame(rows, columns=["symbol", "description", "data_or_model_source", "role"])


def build_calibration_metrics() -> pd.DataFrame:
    frames: list[pd.DataFrame] = []
    sources = [
        (
            "new_must_glycerol_estimability_doe",
            ROOT / "fermentation_model/results/new_must_glycerol_estimability_doe/fit_summary.csv",
        ),
        (
            "secondary_fit_capacity",
            ROOT / "fermentation_model/results/secondary_fit_capacity/fit_capacity_summary.csv",
        ),
    ]
    for workflow, path in sources:
        if not path.exists():
            continue
        df = pd.read_csv(path)
        df.insert(0, "workflow", workflow)
        df.insert(1, "source_file", rel(path))
        frames.append(df)
    if not frames:
        return pd.DataFrame()
    return pd.concat(frames, ignore_index=True, sort=False)


def build_parameter_tables() -> tuple[pd.DataFrame, pd.DataFrame]:
    theta_path = ROOT / "fermentation_model/results/new_must_glycerol_estimability_doe/theta_fit_table.csv"
    theta = pd.read_csv(theta_path) if theta_path.exists() else pd.DataFrame()
    if not theta.empty:
        theta.insert(0, "source_file", rel(theta_path))

    rows: list[dict[str, object]] = []
    for workflow in ["aroma_campaign_doe", "aroma_joint_campaign_doe"]:
        meta_path = ROOT / f"fermentation_model/results/{workflow}/aroma_campaign_metadata.json"
        if not meta_path.exists():
            continue
        meta = json.loads(meta_path.read_text(encoding="utf-8"))
        for species, pars in meta.get("partition_model", {}).items():
            row = {"workflow": workflow, "species": species, "source_file": rel(meta_path)}
            row.update(pars)
            rows.append(row)
    partition = pd.DataFrame(rows)
    return theta, partition


def build_mass_closure_table() -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for workflow in ["aroma_campaign_doe", "aroma_joint_campaign_doe"]:
        sim_path = ROOT / f"fermentation_model/results/{workflow}/aroma_campaign_simulations.csv"
        meta_path = ROOT / f"fermentation_model/results/{workflow}/aroma_campaign_metadata.json"
        if not sim_path.exists() or not meta_path.exists():
            continue
        df = pd.read_csv(sim_path)
        meta = json.loads(meta_path.read_text(encoding="utf-8"))
        eta_lookup = {
            species: float(pars.get("trap_efficiency", np.nan))
            for species, pars in meta.get("partition_model", {}).items()
        }
        for (candidate, species), group in df.groupby(["candidate", "species"], sort=True):
            group = group.sort_values("t")
            t = group["t"].to_numpy(dtype=float)
            synthesis_integral = float(np.trapz(group["A_synthesis_rate"].to_numpy(dtype=float), t))
            loss_integral = float(np.trapz(group["A_loss_rate"].to_numpy(dtype=float), t))
            first = group.iloc[0]
            last = group.iloc[-1]
            liquid_delta = float(last["A_liq"] - first["A_liq"])
            loss_delta = float(last["A_loss"] - first["A_loss"])
            state_total = liquid_delta + loss_delta
            closure_error = synthesis_integral - state_total
            closure_rel = closure_error / max(abs(synthesis_integral), abs(state_total), 1e-12)
            loss_error = loss_integral - loss_delta
            loss_rel = loss_error / max(abs(loss_integral), abs(loss_delta), 1e-12)
            eta = eta_lookup.get(str(species), np.nan)
            cond_expected = float(eta * last["A_loss"]) if np.isfinite(eta) else np.nan
            cond_error = float(last["A_cond_final"] - cond_expected) if np.isfinite(cond_expected) else np.nan
            rows.append(
                {
                    "workflow": workflow,
                    "candidate": candidate,
                    "species": species,
                    "family": last["family"],
                    "t_final_h": float(last["t"]),
                    "liquid_final": float(last["A_liq"]),
                    "loss_final": float(last["A_loss"]),
                    "condensate_final": float(last["A_cond_final"]),
                    "trap_efficiency": eta,
                    "condensate_expected_from_loss": cond_expected,
                    "condensate_abs_error": cond_error,
                    "synthesis_integral_trapz": synthesis_integral,
                    "liquid_plus_loss_delta": state_total,
                    "mass_closure_abs_error": closure_error,
                    "mass_closure_rel_error": closure_rel,
                    "loss_integral_trapz": loss_integral,
                    "loss_state_delta": loss_delta,
                    "loss_balance_abs_error": loss_error,
                    "loss_balance_rel_error": loss_rel,
                    "termination": last.get("termination", ""),
                    "source_file": rel(sim_path),
                }
            )
    return pd.DataFrame(rows)


def build_sensitivity_table() -> pd.DataFrame:
    frames: list[pd.DataFrame] = []
    sources = [
        (
            "new_must_estimability",
            ROOT / "fermentation_model/results/new_must_glycerol_estimability_doe/parameter_estimability_summary.csv",
        ),
        (
            "aroma_synthesis_campaign",
            ROOT / "fermentation_model/results/aroma_campaign_doe/aroma_campaign_parameter_reduction.csv",
        ),
        (
            "aroma_joint_campaign",
            ROOT / "fermentation_model/results/aroma_joint_campaign_doe/aroma_campaign_parameter_reduction.csv",
        ),
        (
            "aroma_joint_eigen",
            ROOT / "fermentation_model/results/aroma_joint_campaign_doe/aroma_eigen_summary.csv",
        ),
    ]
    for workflow, path in sources:
        if path.exists():
            df = pd.read_csv(path)
            df.insert(0, "workflow", workflow)
            df.insert(1, "source_file", rel(path))
            frames.append(df)
    if not frames:
        return pd.DataFrame()
    return pd.concat(frames, ignore_index=True, sort=False)


def make_markdown_table(df: pd.DataFrame, max_rows: int = 12) -> str:
    if df.empty:
        return "_No hay datos disponibles._"
    return df.head(max_rows).to_markdown(index=False)


def build_equations_doc() -> str:
    refs = {
        "base_X": find_line("fermentation_model/run_new_must_glycerol_estimability_doe.py", "def X_balance"),
        "base_N": find_line("fermentation_model/run_new_must_glycerol_estimability_doe.py", "def N_balance"),
        "base_G": find_line("fermentation_model/run_new_must_glycerol_estimability_doe.py", "def G_balance"),
        "base_F": find_line("fermentation_model/run_new_must_glycerol_estimability_doe.py", "def F_balance"),
        "base_E": find_line("fermentation_model/run_new_must_glycerol_estimability_doe.py", "def E_balance"),
        "base_Gly": find_line("fermentation_model/run_new_must_glycerol_estimability_doe.py", "def Gly_balance"),
        "aroma_layer": find_line("fermentation_model/run_aroma_campaign_doe.py", "def add_aroma_layer"),
        "aroma_loss": find_line("fermentation_model/run_aroma_campaign_doe.py", "m.aroma_loss_rate"),
        "aroma_liq_balance": find_line("fermentation_model/run_aroma_campaign_doe.py", "def aroma_liquid_balance"),
        "aroma_loss_balance": find_line("fermentation_model/run_aroma_campaign_doe.py", "def aroma_loss_balance"),
        "aroma_cond": find_line("fermentation_model/run_aroma_campaign_doe.py", "def aroma_condensate_measurement"),
        "unifac": find_line("fermentation_model/aroma_partition_unifac.py", "def unifac_partition_K"),
    }
    return f"""# Ecuaciones y supuestos tecnicos - A02

Actividad: {ACTIVITY_CODE} - {ACTIVITY_NAME}

## Modelo base de fermentacion

El estado dinamico base utilizado para calibracion y diseno experimental es:

`X, Xd, N, G, F, E, Gly`

con:

- `X`: biomasa viable.
- `Xd`: biomasa muerta.
- `N`: nitrogeno asimilable.
- `G`: glucosa.
- `F`: fructosa.
- `E`: etanol.
- `Gly`: glicerol.

Las ecuaciones estan implementadas en `fermentation_model/run_new_must_glycerol_estimability_doe.py`:

- `X_balance`: linea {refs['base_X']}.
- `N_balance`: linea {refs['base_N']}.
- `G_balance`: linea {refs['base_G']}.
- `F_balance`: linea {refs['base_F']}.
- `E_balance`: linea {refs['base_E']}.
- `Gly_balance`: linea {refs['base_Gly']}.

Forma resumida:

```text
dX/dt   = (mu - kd) X + X_input
dXd/dt  = kd X
dN/dt   = -qN phi_N X + N_input
dG/dt   = -(qXG phi_N + qEG phi_G + m_G) X + G_input
dF/dt   = -(qXF phi_N + qEF phi_F + m_F) X + F_input
dE/dt   = (betaG phi_G + betaF phi_F) X + E_input
dGly/dt = (gammaG0 phi_G + gammaF0 phi_F) X
```

Los factores `phi_G` y `phi_F` incorporan temperatura, saturacion de sustrato, interaccion glucosa/fructosa e inhibicion por etanol.

## Capa de perdidas aromaticas

La capa aromatica esta implementada en `fermentation_model/run_aroma_campaign_doe.py`, desde `add_aroma_layer` en linea {refs['aroma_layer']}.

Especies incluidas:

- `ethyl_acetate`
- `isoamyl_acetate`
- `ethyl_octanoate`

Forma resumida:

```text
Q_CO2(t)        = conversion_CO2 * CO2_rate(t)
K_lg(i,t)       = K20_i exp(aT_i (T-20) + aE_i (E-50) + aS_i (G+F-100))
r_syn(i,t)      = (k_growth_i phi_growth + k_stationary_i (1-phi_growth)) sugar_uptake(t)
r_loss(i,t)     = alpha_i K_lg(i,t) Q_CO2(t) A_liq(i,t)
dA_liq(i,t)/dt  = r_syn(i,t) - r_loss(i,t)
dA_loss(i,t)/dt = r_loss(i,t)
A_cond(i)       = eta_i A_loss(i,t_final)
```

Referencias de codigo:

- `aroma_loss_rate`: linea {refs['aroma_loss']}.
- `aroma_liquid_balance`: linea {refs['aroma_liq_balance']}.
- `aroma_loss_balance`: linea {refs['aroma_loss_balance']}.
- `aroma_condensate_measurement`: linea {refs['aroma_cond']}.

## Particion gas-liquido

El coeficiente `K_lg` proviene de una regresion log-lineal ajustada sobre calculos UNIFAC de dilucion infinita. La funcion base `unifac_partition_K` esta en `fermentation_model/aroma_partition_unifac.py`, linea {refs['unifac']}.

El modo operacional usa:

```text
log(K_lg) = log(K20) + temp_slope (T - 20) + ethanol_slope (E - 50) + sugar_slope (G + F - 100)
```

## Supuestos tecnicos principales

- La fructosa se aproxima como glucosa-equivalente en la composicion liquida UNIFAC cuando no existe asignacion directa disponible.
- Las perdidas aromaticas se calculan por stripping proporcional a `Q_CO2`, `K_lg` y concentracion liquida `A_liq`.
- Las eficiencias de trampa terminal son fijas por especie en esta version del modelo.
- Los parametros de particion se fijan por defecto para evitar confundir equilibrio gas-liquido con eficiencia de trampa.
- El condensado terminal aporta cierre de masa acumulado, no dinamica temporal de gas.
- La calibracion base de fermentacion usa datos VL3/natural/sintetico; la capa aromatica queda documentada como extension mecanistica con DOE/FIM y cierre de masa simulado.
"""


def build_compliance_matrix() -> pd.DataFrame:
    rows = [
        {
            "evidencia_minima_requerida": "Ecuaciones y supuestos del balance",
            "estado_bundle": "cubierto",
            "archivos_bundle": "03_ecuaciones_y_supuestos_A02.md; code_sources/*.py",
            "detalle": "Incluye balances base X/Xd/N/G/F/E/Gly, capa A_liq/A_loss/A_cond y supuestos UNIFAC/CO2.",
        },
        {
            "evidencia_minima_requerida": "Datos de calibracion",
            "estado_bundle": "cubierto",
            "archivos_bundle": "data_sources/*.xlsx; tables/02_resumen_datos_fuente.csv; tables/02_diccionario_variables_clave.csv",
            "detalle": "Incluye libros VL3 y mosto sintetico, mas resumen de hojas, columnas y variables.",
        },
        {
            "evidencia_minima_requerida": "Parametros",
            "estado_bundle": "cubierto",
            "archivos_bundle": "tables/05_parametros_cineticos.csv; tables/05_parametros_particion_aromas.csv",
            "detalle": "Incluye theta calibrado y parametros de particion aromaticos derivados de UNIFAC.",
        },
        {
            "evidencia_minima_requerida": "Cierre de masa",
            "estado_bundle": "cubierto",
            "archivos_bundle": "tables/06_cierre_masa_aromatico.csv; 03_ecuaciones_y_supuestos_A02.md",
            "detalle": "Cierre por especie/candidato: integral de sintesis, liquido final, perdida final y condensado esperado.",
        },
        {
            "evidencia_minima_requerida": "Metricas de ajuste",
            "estado_bundle": "cubierto",
            "archivos_bundle": "tables/04_metricas_calibracion.csv; source_reports/*fit*; source_reports/*glycerol*",
            "detalle": "Incluye WSSE, residuales, grados de libertad, exito de ajuste y capacidad de ajuste secundaria.",
        },
        {
            "evidencia_minima_requerida": "Version de codigo",
            "estado_bundle": "cubierto",
            "archivos_bundle": "08_version_codigo.md; 00_manifest.csv",
            "detalle": "Incluye commit, rama, estado git, diff stat y hashes SHA256 de artefactos.",
        },
        {
            "evidencia_minima_requerida": "Analisis de sensibilidad",
            "estado_bundle": "cubierto",
            "archivos_bundle": "tables/07_sensibilidad_estimabilidad.csv; source_reports/*eigen*; source_reports/*campaign*",
            "detalle": "Incluye FIM, reduccion de varianza, ranking DOE, eigenvalores y direcciones debiles.",
        },
    ]
    return pd.DataFrame(rows)


def build_proposed_row() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "Objetivo": "OE2",
                "Actividad": "2.13",
                "Estado Gantt": "Nueva",
                "Nombre actividad (fuente)": ACTIVITY_NAME,
                "Inicio Gantt": "2025-05-01",
                "Termino Gantt": "2025-09-30",
                "Alcance provisional IA5": "Si",
                "Prioridad": "Alta",
                "Evidencia minima requerida": "Ecuaciones y supuestos del balance, datos de calibracion, parametros, cierre de masa, metricas de ajuste, version de codigo y analisis de sensibilidad.",
                "Anexo propuesto": "A02",
                "Referentes segun Gantt RRHH": "Cristobal Torrealba; Ricardo Luna; Sergio Manzano; Mauricio Araya; equipo analitico",
                "Estado evidencia": "Completa en bundle A02 para cobertura documental de evidencia minima",
                "% avance tecnico": "100% documental; validacion experimental directa de condensado/gas queda trazada si se exige como continuidad",
                "Inicio real": "2025-05-01, segun periodo experimental/Gantt",
                "Termino real": f"{GENERATED_DATE}, cierre documental del bundle tecnico",
                "Desfase/modificacion": "Si: consolidacion computacional y empaquetamiento documental posterior al termino Gantt; datos experimentales asociados a campanas 2025.",
                "Resumen de resultado para informe": "Se desarrollo y calibro un modelo dinamico de balance de masa para fermentacion que integra biomasa, nitrogeno, glucosa, fructosa, etanol y glicerol usando datos VL3 y mosto sintetico/natural. Sobre esta base se implemento una extension para perdidas aromaticas por stripping con CO2, particion gas-liquido derivada de UNIFAC, acumulacion de perdidas y prediccion de condensado terminal. El bundle A02 documenta ecuaciones, supuestos, datos fuente, parametros, metricas de ajuste, cierre de masa aromatico, version de codigo y analisis de sensibilidad/estimabilidad.",
                "Ruta/archivo fuente": "rendicion_tecnica/A02_balance_masa_perdidas_aromaticas/README_A02_bundle.md",
                "Responsable de cierre": "Cristobal Torrealba",
                "Alerta/observacion": "Usar '100%' como cobertura documental de evidencia minima. No declarar validacion experimental directa de gas/condensado salvo que se adjunten mediciones externas.",
                "Fuente": "Repositorio pyomo-doe; Carta Gantt Modificada (Solicitud PI-4497)_2025_03_27.xlsx",
            }
        ]
    )


def build_version_doc() -> str:
    key_files = [
        ROOT / "fermentation_model/run_new_must_glycerol_estimability_doe.py",
        ROOT / "fermentation_model/run_aroma_campaign_doe.py",
        ROOT / "fermentation_model/aroma_partition_unifac.py",
        ROOT / "fermentation_model/new_must_data_loader.py",
        ROOT / "fermentation_model/data/Calibration_data_vl3.xlsx",
        ROOT / "fermentation_model/data/mosto_sintetico_vl3.xlsx",
    ]
    hashes = []
    for path in key_files:
        if path.exists():
            hashes.append(f"- `{rel(path)}`: `{sha256(path)}`")
    return f"""# Version de codigo y trazabilidad

Generado: {datetime.now().astimezone().isoformat()}
Fecha de cierre documental usada para la rendicion: {GENERATED_DATE}

## Git

- Rama: `{run_git(['branch', '--show-current'])}`
- Commit HEAD: `{run_git(['rev-parse', 'HEAD'])}`
- Commit corto: `{run_git(['rev-parse', '--short', 'HEAD'])}`

## Estado de trabajo

```text
{run_git(['status', '--short'])}
```

## Diff stat

```text
{run_git(['diff', '--stat'])}
```

## Hashes SHA256 de fuentes clave

{chr(10).join(hashes)}

## Nota de interpretacion

El arbol de trabajo contiene archivos sin seguimiento y archivos modificados. Para rendicion, usar este documento junto con `00_manifest.csv`, que registra hashes de los artefactos copiados al bundle. Si se requiere congelar formalmente la evidencia, crear commit/tag o almacenar el ZIP en el gestor documental del proyecto.
"""


def build_readme(
    data_summary: pd.DataFrame,
    calibration: pd.DataFrame,
    mass_closure: pd.DataFrame,
    sensitivity: pd.DataFrame,
) -> str:
    if mass_closure.empty:
        max_closure = "NA"
        max_loss = "NA"
    else:
        max_closure = f"{mass_closure['mass_closure_rel_error'].abs().max():.4%}"
        max_loss = f"{mass_closure['loss_balance_rel_error'].abs().max():.4%}"
    return f"""# Bundle A02 - Balance de masa y perdidas aromaticas

Actividad: {ACTIVITY_CODE} - {ACTIVITY_NAME}

Fecha de cierre documental: {GENERATED_DATE}

## Proposito

Este bundle organiza la evidencia tecnica para cubrir el 100% de las categorias solicitadas en la fila de rendicion: ecuaciones/supuestos, datos de calibracion, parametros, cierre de masa, metricas de ajuste, version de codigo y analisis de sensibilidad.

La cobertura es documental y tecnica. Si una revision externa exige validacion experimental directa de gas o condensado, usar este paquete como base y anexar esas mediciones como una capa adicional.

## Archivos principales

- `ANEXO_A02_borrador.md`: texto base listo para transformar en anexo.
- `01_matriz_cumplimiento_A02.csv`: mapeo requisito -> evidencia.
- `02_fila_propuesta_planilla_A02.csv`: fila sugerida para la planilla de seguimiento.
- `03_ecuaciones_y_supuestos_A02.md`: ecuaciones, supuestos y referencias a codigo.
- `08_version_codigo.md`: commit, estado git y hashes.
- `00_manifest.csv`: inventario completo del bundle con rutas y SHA256.

## Tablas generadas

- `tables/02_resumen_datos_fuente.csv`: resumen de libros Excel, hojas, columnas y variables.
- `tables/02_diccionario_variables_clave.csv`: diccionario de estados/variables.
- `tables/04_metricas_calibracion.csv`: metricas de ajuste y capacidad de ajuste.
- `tables/05_parametros_cineticos.csv`: parametros cineticos calibrados.
- `tables/05_parametros_particion_aromas.csv`: parametros de particion aromaticos UNIFAC/surrogado.
- `tables/06_cierre_masa_aromatico.csv`: cierre de masa por candidato y especie.
- `tables/07_sensibilidad_estimabilidad.csv`: FIM, estimabilidad, reduccion de varianza y eigen-resumen.

## Resumen cuantitativo

- Filas de resumen de datos fuente: {len(data_summary)}
- Filas de metricas/calibracion: {len(calibration)}
- Filas de cierre de masa aromatico: {len(mass_closure)}
- Error relativo maximo de cierre de masa aromatico por integral numerica: {max_closure}
- Error relativo maximo de balance de perdida aromatico por integral numerica: {max_loss}
- Filas de sensibilidad/estimabilidad: {len(sensitivity)}

## Como usarlo para construir el anexo

1. Abrir `ANEXO_A02_borrador.md` y usarlo como cuerpo narrativo del anexo.
2. Insertar o convertir a tablas formales los CSV de `tables/`.
3. Adjuntar como respaldo los reportes de `source_reports/`, las figuras de `figures/` y los datos en `data_sources/`.
4. Usar `01_matriz_cumplimiento_A02.csv` para demostrar que cada requisito minimo tiene una evidencia concreta.
5. Usar `02_fila_propuesta_planilla_A02.csv` para completar la planilla de control de rendicion.
6. Si se actualizan datos o resultados, regenerar con:

```powershell
python scripts/build_A02_evidence_bundle.py
```
"""


def build_annex_draft(
    compliance: pd.DataFrame,
    data_summary: pd.DataFrame,
    calibration: pd.DataFrame,
    theta: pd.DataFrame,
    partition: pd.DataFrame,
    mass_closure: pd.DataFrame,
    sensitivity: pd.DataFrame,
) -> str:
    selected_calibration_cols = [
        col
        for col in [
            "workflow",
            "fit",
            "fit_label",
            "n_batches",
            "mediums",
            "n_parameters",
            "success",
            "initial_wsse",
            "final_wsse",
            "n_residuals",
            "wsse_per_residual",
            "wsse_per_dof",
        ]
        if col in calibration.columns
    ]
    calibration_preview = calibration[selected_calibration_cols] if selected_calibration_cols else calibration
    theta_preview = theta.head(6) if not theta.empty else theta
    partition_cols = [
        col
        for col in [
            "workflow",
            "species",
            "K20",
            "temp_slope",
            "ethanol_slope",
            "sugar_slope",
            "fit_rmse_log",
            "fit_max_abs_log_error",
            "trap_efficiency",
        ]
        if col in partition.columns
    ]
    partition_preview = partition[partition_cols] if partition_cols else partition
    closure_cols = [
        "workflow",
        "candidate",
        "species",
        "liquid_final",
        "loss_final",
        "condensate_final",
        "mass_closure_rel_error",
        "loss_balance_rel_error",
    ]
    closure_preview = mass_closure[closure_cols] if not mass_closure.empty else mass_closure
    return f"""# Anexo A02 - Desarrollo y calibracion de modelos de balance de masa para cuantificar perdidas aromaticas

## Identificacion

- Objetivo especifico: OE2
- Actividad: 2.13
- Nombre actividad: {ACTIVITY_NAME}
- Periodo Gantt: 2025-05-01 a 2025-09-30
- Fecha de cierre documental del bundle: {GENERATED_DATE}
- Estado propuesto de evidencia: cobertura documental completa de la evidencia minima requerida.

## Resumen ejecutivo

Se desarrollo y calibro un modelo dinamico de balance de masa para fermentacion que integra biomasa, nitrogeno asimilable, glucosa, fructosa, etanol y glicerol. El modelo fue parametrizado y evaluado usando bases VL3, mosto sintetico y datos natural/sintetico homologados. Sobre esta base se implemento una capa mecanistica de perdidas aromaticas por stripping con CO2, con particion gas-liquido derivada de calculos UNIFAC y una prediccion de condensado terminal para cierre de masa.

El paquete A02 cubre las categorias solicitadas: ecuaciones y supuestos, datos de calibracion, parametros, cierre de masa, metricas de ajuste, version de codigo y analisis de sensibilidad/estimabilidad.

## Matriz de cumplimiento

{make_markdown_table(compliance, max_rows=20)}

## Datos de calibracion

Los datos fuente principales son:

- `Calibration_data_vl3.xlsx`: fermentaciones historicas VL3 con azucares, nitrogeno, glicerol, piruvato, acetaldehido, etanol y aromas totales.
- `mosto_sintetico_vl3.xlsx`: datos homologados de mosto sintetico, diseno CCD, datos originales, flags de calidad y diccionario de mapeo.
- `mosto_natural_xthiol.xlsx`: fuente complementaria natural usada por el flujo de carga cuando esta disponible.

Resumen de hojas:

{make_markdown_table(data_summary[['source_file', 'sheet', 'data_rows_estimated', 'columns', 'aroma_total_columns']], max_rows=18) if not data_summary.empty else '_No hay resumen de datos._'}

## Modelo y supuestos

El detalle formal esta en `03_ecuaciones_y_supuestos_A02.md`. En sintesis, el modelo base resuelve balances de masa para `X, Xd, N, G, F, E, Gly`; la extension aromatica agrega estados `A_liq`, `A_loss` y `A_cond` para cuantificar aroma retenido, aroma perdido y condensado terminal.

La perdida aromatica se calcula como:

```text
r_loss(i,t) = alpha_i K_lg(i,t) Q_CO2(t) A_liq(i,t)
```

y el cierre acumulado queda:

```text
dA_liq/dt  = r_syn - r_loss
dA_loss/dt = r_loss
A_cond     = eta A_loss(t_final)
```

## Parametros y calibracion

Metricas principales de calibracion:

{make_markdown_table(calibration_preview, max_rows=12)}

Parametros cineticos:

{make_markdown_table(theta_preview, max_rows=8)}

Parametros de particion aromatica:

{make_markdown_table(partition_preview, max_rows=8)}

## Cierre de masa aromatico

La tabla `tables/06_cierre_masa_aromatico.csv` resume el cierre por candidato y especie. La columna `mass_closure_rel_error` compara la integral numerica de sintesis con el incremento de `A_liq + A_loss`. La columna `loss_balance_rel_error` compara la integral numerica de perdida con el incremento de `A_loss`. Estos errores son diagnosticos de postproceso numerico; las ecuaciones del modelo imponen el cierre en la discretizacion.

Vista preliminar:

{make_markdown_table(closure_preview, max_rows=12)}

## Sensibilidad y estimabilidad

El paquete incluye analisis FIM, reduccion de varianza, ranking DOE y diagnostico de eigen-direcciones. La tabla consolidada esta en `tables/07_sensibilidad_estimabilidad.csv`.

Vista preliminar:

{make_markdown_table(sensitivity, max_rows=12)}

## Archivos de respaldo

- Datos fuente: `data_sources/`
- Codigo fuente clave: `code_sources/`
- Notebooks de trabajo: `notebooks/`
- Reportes originales generados por scripts: `source_reports/`
- Tablas originales de resultados: `source_tables/`
- Figuras de campanas y diagnosticos: `figures/`
- Inventario con hashes: `00_manifest.csv`

## Texto sugerido para informe

Se desarrollo y calibro un modelo dinamico de balance de masa para fermentacion que integra biomasa, nitrogeno, glucosa, fructosa, etanol y glicerol usando datos VL3 y mosto sintetico/natural. Sobre esta base se implemento una extension para perdidas aromaticas por stripping con CO2, particion gas-liquido derivada de UNIFAC, acumulacion de perdidas y prediccion de condensado terminal. El bundle A02 documenta ecuaciones, supuestos, datos fuente, parametros, metricas de ajuste, cierre de masa aromatico, version de codigo y analisis de sensibilidad/estimabilidad, cubriendo la evidencia minima requerida para la actividad OE2-2.13.

## Limitacion declarable

La evidencia cubre el desarrollo, calibracion base, simulacion y diseno experimental de perdidas aromaticas. Si la rendicion exige validacion experimental directa de gas o condensado, debe anexarse la medicion externa correspondiente como complemento; el presente bundle deja el punto trazado y preparado mediante la tabla de cierre de masa y el balance `A_cond = eta A_loss(t_final)`.
"""


def build_manifest_markdown(manifest: pd.DataFrame) -> str:
    counts = manifest.groupby(["category", "status"]).size().reset_index(name="n")
    return f"""# Manifest A02

Este archivo resume el inventario completo del bundle. La version tabular con hashes esta en `00_manifest.csv`.

## Conteo por categoria

{make_markdown_table(counts, max_rows=100)}

## Primeras rutas

{make_markdown_table(manifest[['category', 'evidence_requirement', 'status', 'bundle_path', 'source_path']], max_rows=30)}
"""


def copy_source_artifacts() -> None:
    data_sources = [
        "fermentation_model/data/Calibration_data_vl3.xlsx",
        "fermentation_model/data/mosto_sintetico_vl3.xlsx",
        "fermentation_model/data/mosto_natural_xthiol.xlsx",
    ]
    for source in data_sources:
        copy_file(source, "data_sources", "datos", "Datos de calibracion", "Libro de datos fuente")

    code_sources = [
        "fermentation_model/run_new_must_glycerol_estimability_doe.py",
        "fermentation_model/run_aroma_campaign_doe.py",
        "fermentation_model/aroma_partition_unifac.py",
        "fermentation_model/new_must_data_loader.py",
        "fermentation_model/run_secondary_fit_capacity.py",
        "fermentation_model/run_secondary_joint_campaign_doe.py",
        "scripts/build_A02_evidence_bundle.py",
    ]
    for source in code_sources:
        copy_file(source, "code_sources", "codigo", "Version de codigo", "Codigo fuente clave")

    notebooks = [
        "fermentation_model/fermentation_model_calibration.ipynb",
        "fermentation_model/fermentation_model_calibration_3_effective_reformulation.ipynb",
        "fermentation_model/fermentation_new_must_glycerol_estimability_doe.executed.ipynb",
        "fermentation_model/fermentation_secondary_fit_capacity.executed.ipynb",
        "fermentation_model/fermentation_secondary_joint_campaign_doe.executed.ipynb",
        "fermentation_model/fermentation_aroma_campaign_doe.ipynb",
    ]
    for source in notebooks:
        copy_file(source, "notebooks", "notebooks", "Metricas de ajuste y analisis", "Notebook de respaldo")

    reports = [
        "fermentation_model/results/new_must_glycerol_estimability_doe/new_must_glycerol_estimability_doe_report.md",
        "fermentation_model/results/secondary_fit_capacity/secondary_fit_capacity_report.md",
        "fermentation_model/results/secondary_joint_campaign_doe/secondary_joint_campaign_report.md",
        "fermentation_model/results/aroma_campaign_doe/aroma_campaign_report.md",
        "fermentation_model/results/aroma_joint_campaign_doe/aroma_campaign_report.md",
        "fermentation_model/results/aroma_joint_campaign_doe/aroma_eigen_analysis_report.md",
        "fermentation_model/results/aroma_joint_campaign_doe/aroma_dopt_benchmark_report.md",
        "fermentation_model/results/aroma_symbolic_unifac_trial_report.md",
    ]
    for source in reports:
        copy_file(
            source,
            "source_reports",
            "reportes",
            "Metricas, sensibilidad y supuestos",
            "Reporte fuente",
            namespace=True,
        )

    source_tables = [
        "fermentation_model/results/new_must_glycerol_estimability_doe/fit_summary.csv",
        "fermentation_model/results/new_must_glycerol_estimability_doe/theta_fit_table.csv",
        "fermentation_model/results/new_must_glycerol_estimability_doe/parameter_estimability_summary.csv",
        "fermentation_model/results/new_must_glycerol_estimability_doe/profile_summary_combined.csv",
        "fermentation_model/results/secondary_fit_capacity/fit_capacity_summary.csv",
        "fermentation_model/results/secondary_fit_capacity/fit_capacity_state_metrics.csv",
        "fermentation_model/results/aroma_campaign_doe/aroma_campaign_parameter_reduction.csv",
        "fermentation_model/results/aroma_campaign_doe/aroma_campaign_selected.csv",
        "fermentation_model/results/aroma_campaign_doe/aroma_candidate_ranking.csv",
        "fermentation_model/results/aroma_campaign_doe/aroma_campaign_simulations.csv",
        "fermentation_model/results/aroma_campaign_doe/aroma_campaign_metadata.json",
        "fermentation_model/results/aroma_joint_campaign_doe/aroma_campaign_parameter_reduction.csv",
        "fermentation_model/results/aroma_joint_campaign_doe/aroma_campaign_selected.csv",
        "fermentation_model/results/aroma_joint_campaign_doe/aroma_candidate_ranking.csv",
        "fermentation_model/results/aroma_joint_campaign_doe/aroma_campaign_simulations.csv",
        "fermentation_model/results/aroma_joint_campaign_doe/aroma_campaign_metadata.json",
        "fermentation_model/results/aroma_joint_campaign_doe/aroma_eigen_summary.csv",
        "fermentation_model/results/aroma_joint_campaign_doe/aroma_eigen_spectrum.csv",
        "fermentation_model/results/aroma_joint_campaign_doe/aroma_weak_eigendirection_top_loadings.csv",
    ]
    for source in source_tables:
        copy_file(
            source,
            "source_tables",
            "tablas_fuente",
            "Parametros, cierre, metricas y sensibilidad",
            "Tabla fuente",
            namespace=True,
        )

    copy_glob(
        "fermentation_model/results/aroma_campaign_doe/*.png",
        "figures",
        "figuras",
        "Analisis de sensibilidad y diseno",
        "Figura campana aromas",
    )
    copy_glob(
        "fermentation_model/results/aroma_joint_campaign_doe/*.png",
        "figures",
        "figuras",
        "Analisis de sensibilidad y diseno",
        "Figura campana conjunta aromas",
    )
    copy_glob(
        "fermentation_model/results/new_must_glycerol_estimability_doe/plots/*.png",
        "figures",
        "figuras",
        "Datos, calibracion y DOE",
        "Figura nueva base mosto/glicerol",
    )


def write_zip() -> None:
    if ZIP_PATH.exists():
        ZIP_PATH.unlink()
    with zipfile.ZipFile(ZIP_PATH, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for path in sorted(BUNDLE_DIR.rglob("*")):
            if path.is_file():
                zf.write(path, path.relative_to(BUNDLE_DIR.parent))


def main() -> int:
    reset_bundle()
    copy_source_artifacts()

    data_summary = build_data_summary()
    variable_dict = build_variable_dictionary()
    calibration = build_calibration_metrics()
    theta, partition = build_parameter_tables()
    mass_closure = build_mass_closure_table()
    sensitivity = build_sensitivity_table()
    compliance = build_compliance_matrix()
    proposed_row = build_proposed_row()

    write_df_artifact(
        "01_matriz_cumplimiento_A02.csv",
        compliance,
        "matriz_cumplimiento",
        "Evidencia minima requerida",
        "Matriz requisito a evidencia",
    )
    write_df_artifact(
        "02_fila_propuesta_planilla_A02.csv",
        proposed_row,
        "planilla",
        "Estado evidencia y resumen para informe",
        "Fila sugerida para planilla de control",
    )
    write_df_artifact(
        "tables/02_resumen_datos_fuente.csv",
        data_summary,
        "datos",
        "Datos de calibracion",
        "Resumen de libros Excel y hojas",
    )
    write_df_artifact(
        "tables/02_diccionario_variables_clave.csv",
        variable_dict,
        "datos",
        "Ecuaciones y supuestos del balance",
        "Diccionario de variables y estados",
    )
    write_df_artifact(
        "tables/04_metricas_calibracion.csv",
        calibration,
        "metricas",
        "Metricas de ajuste",
        "Metricas consolidadas de calibracion",
    )
    write_df_artifact(
        "tables/05_parametros_cineticos.csv",
        theta,
        "parametros",
        "Parametros",
        "Parametros cineticos calibrados",
    )
    write_df_artifact(
        "tables/05_parametros_particion_aromas.csv",
        partition,
        "parametros",
        "Parametros",
        "Parametros de particion aromaticos",
    )
    write_df_artifact(
        "tables/06_cierre_masa_aromatico.csv",
        mass_closure,
        "cierre_masa",
        "Cierre de masa",
        "Cierre de masa aromatico por candidato y especie",
    )
    write_df_artifact(
        "tables/07_sensibilidad_estimabilidad.csv",
        sensitivity,
        "sensibilidad",
        "Analisis de sensibilidad",
        "Sensibilidad, estimabilidad y reduccion de varianza",
    )

    equations_doc = build_equations_doc()
    write_text_artifact(
        "03_ecuaciones_y_supuestos_A02.md",
        equations_doc,
        "ecuaciones",
        "Ecuaciones y supuestos del balance",
        "Documento tecnico de ecuaciones y supuestos",
    )
    write_text_artifact(
        "08_version_codigo.md",
        build_version_doc(),
        "version_codigo",
        "Version de codigo",
        "Trazabilidad de version de codigo",
    )
    write_text_artifact(
        "README_A02_bundle.md",
        build_readme(data_summary, calibration, mass_closure, sensitivity),
        "guia",
        "Uso del bundle",
        "Guia principal de uso del paquete",
    )
    write_text_artifact(
        "ANEXO_A02_borrador.md",
        build_annex_draft(compliance, data_summary, calibration, theta, partition, mass_closure, sensitivity),
        "anexo",
        "Anexo propuesto",
        "Borrador narrativo del anexo A02",
    )

    manifest_for_markdown = pd.DataFrame(manifest_rows)
    manifest_md_path = BUNDLE_DIR / "00_manifest.md"
    manifest_md_path.write_text(build_manifest_markdown(manifest_for_markdown), encoding="utf-8")
    record_artifact(
        bundle_path=manifest_md_path,
        category="manifest",
        evidence_requirement="Inventario y trazabilidad",
        role="Resumen legible del manifest",
    )

    manifest_path = BUNDLE_DIR / "00_manifest.csv"
    manifest_rows.append(
        {
            "activity": ACTIVITY_CODE,
            "category": "manifest",
            "evidence_requirement": "Inventario y trazabilidad",
            "role_in_bundle": "Inventario completo con hashes; fila autoreferencial sin hash estable",
            "status": "self_index",
            "source_path": "generated",
            "bundle_path": rel(manifest_path),
            "bytes": "",
            "sha256": "",
        }
    )
    manifest = pd.DataFrame(manifest_rows)
    manifest.to_csv(manifest_path, index=False, quoting=csv.QUOTE_MINIMAL)

    write_zip()
    print(f"Bundle written: {BUNDLE_DIR}")
    print(f"ZIP written: {ZIP_PATH}")
    print(f"Artifacts indexed: {len(manifest_rows)}")
    if not mass_closure.empty:
        print(f"Max aroma mass closure rel error: {mass_closure['mass_closure_rel_error'].abs().max():.6f}")
        print(f"Max aroma loss balance rel error: {mass_closure['loss_balance_rel_error'].abs().max():.6f}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
